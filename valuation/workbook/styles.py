from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

NUM_FMT = "#,##0.00"
PCT_FMT = "0.00%"

STRUCTURE_STYLES = {
    "sheet_title": {"fill": "1F4E78", "font": "FFFFFF"},
    "section": {"fill": "ED7D31", "font": "FFFFFF"},
    "group": {"fill": "FFF2CC", "font": "1F1F1F"},
    "table_header": {"fill": "D9EAF7", "font": "1F4E78"},
    "footnote": {"fill": None, "font": "888888"},
}

SHEET_TITLE_SUFFIX = {
    "总结": "估值模型总结",
    "主营业务拆分": "主营业务拆分",
    "收入预测": "收入预测",
    "运营成本预测": "运营成本预测",
    "利润表预测": "利润表预测",
    "估值预测": "估值预测",
    "历史财务数据": "财务报表",
}

FIRST_SECTION_TITLE = {
    "收入预测": "收入驱动假设区",
    "运营成本预测": "成本费用假设区",
    "利润表预测": "利润表预测",
    "估值预测": "机构一致性预期",
    "历史财务数据": "利润表（亿元）",
}

THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def apply_structure_style(ws: Worksheet, row_cells, style_key: str) -> None:
    if style_key == "group" and ws.title != "收入预测":
        raise ValueError("黄色 group 样式仅用于收入预测的分部名称（预测方法）行")
    style = STRUCTURE_STYLES[style_key]
    fill = style["fill"]
    font_color = style["font"]
    for cell in row_cells:
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(
            color=font_color,
            bold=style_key != "footnote",
            italic=style_key == "footnote",
        )


def write_sheet_title(ws: Worksheet, company_name: str, last_col: int) -> None:
    expected_title = f"{company_name}-{SHEET_TITLE_SUFFIX[ws.title]}"
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row == merged.max_row == 1:
            ws.unmerge_cells(str(merged))
    ws["A1"] = expected_title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    apply_structure_style(ws, [ws["A1"]], "sheet_title")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22


def write_first_section_header(ws: Worksheet, headers_after_a: list[str], row: int = 2) -> None:
    section_title = FIRST_SECTION_TITLE[ws.title]
    values = [section_title, *headers_after_a]
    cells = [ws.cell(row=row, column=col, value=value) for col, value in enumerate(values, 1)]
    apply_structure_style(ws, cells, "section")
    for cell in cells:
        cell.alignment = Alignment(horizontal="center", vertical="center")


def write_section_row(ws: Worksheet, row: int, values: list, last_col: int) -> None:
    cells = []
    for col in range(1, last_col + 1):
        value = values[col - 1] if col - 1 < len(values) else None
        cell = ws.cell(row=row, column=col, value=value)
        cells.append(cell)
    apply_structure_style(ws, cells, "section")
    for cell in cells:
        cell.alignment = Alignment(horizontal="center", vertical="center")


def write_group_row(ws: Worksheet, row: int, label: str, last_col: int) -> None:
    cells = [ws.cell(row=row, column=1, value=label)]
    for col in range(2, last_col + 1):
        cells.append(ws.cell(row=row, column=col, value=None))
    apply_structure_style(ws, cells, "group")


def set_col_widths(ws: Worksheet, last_col: int, first: float = 28.0, other: float = 12.0) -> None:
    ws.column_dimensions["A"].width = first
    for col in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = other


def apply_number_formats(ws: Worksheet, ratio_rows: set[int], last_col: int, note_col: int | None) -> None:
    for row in range(2, ws.max_row + 1):
        label = str(ws.cell(row=row, column=1).value or "")
        for col in range(2, last_col + 1):
            cell = ws.cell(row=row, column=col)
            if cell.value is None:
                continue
            if isinstance(cell.value, str) and not str(cell.value).startswith("="):
                continue
            if cell.number_format == "yyyy-mm-dd":
                continue
            if row in ratio_rows or any(k in label for k in ("率", "增速", "占比", "空间")):
                if "PE" in label or "股本" in label:
                    cell.number_format = NUM_FMT
                else:
                    cell.number_format = PCT_FMT
            else:
                cell.number_format = NUM_FMT
