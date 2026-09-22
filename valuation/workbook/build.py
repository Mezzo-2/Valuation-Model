from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from valuation.segment_research.methods import (
    RATE_FIELDS,
    as_growth_rate,
    assume_unit,
    build_unit,
    fx_unit_note,
    method_spec,
)
from valuation.workbook.anchors import (
    cell_ref,
    cell_ref_at,
    find_col_by_header,
    find_row_by_label,
    write_note,
)
from valuation.historical_financials.company_info import (
    market_rows as company_market_rows,
    profile_rows as company_profile_rows,
)
from valuation.historical_financials.subjects import (
    BALANCE_BOLD,
    BALANCE_ORDER,
    FS_LEFT_COL,
    FS_RIGHT_COL,
    HIST_METRICS,
    INCOME_BOLD,
    INCOME_ORDER,
    MARKET_ROWS,
    fs_label,
)
from valuation.shared.io import artifacts_dir, spec_dir, workbook_filename
from valuation.workbook.load_json import load_spec
from valuation.workbook.tree import (
    build_split_groups,
    segment_group_title,
    top_level_names,
)
from valuation.workbook.styles import (
    NUM_FMT,
    PCT_FMT,
    apply_number_formats,
    apply_structure_style,
    set_col_widths,
    write_first_section_header,
    write_group_row,
    write_section_row,
    write_sheet_title,
)

SHEET_ORDER = [
    "总结",
    "主营业务拆分",
    "收入预测",
    "运营成本预测",
    "利润表预测",
    "估值预测",
    "历史财务数据",
]

def build_workbook(run_dir: Path) -> Path:
    try:
        spec = load_spec(run_dir, require_notes=True)
    except FileNotFoundError as exc:
        raise SystemExit(str(exc)) from exc
    facts = spec["facts"]
    hist = facts["hist_periods"]
    fcst = facts["forecast_periods"]
    periods = hist + fcst

    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for name in SHEET_ORDER:
        wb.create_sheet(name)

    _hist(wb, spec)
    _split(wb, spec)
    _revenue(wb, spec)
    _cost(wb, spec)
    _pnl(wb, spec)
    _valuation(wb, spec)
    _summary(wb, spec)

    wb._sheets.sort(key=lambda ws: SHEET_ORDER.index(ws.title))
    out = artifacts_dir(run_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / workbook_filename(facts.get("company"), facts.get("ticker"))
    wb.save(path)
    return path


def build_hist_workbook(run_dir: Path) -> Path:
    from valuation.shared.io import load_json

    facts = load_json(spec_dir(run_dir) / "facts.json")
    wb = Workbook()
    default = wb.active
    default.title = "历史财务数据"
    _hist(wb, {"facts": facts})
    out = artifacts_dir(run_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / "historical_financials.xlsx"
    wb.save(path)
    return path


def _hist(wb: Workbook, spec: dict) -> None:
    facts = spec["facts"]
    hist = facts["hist_periods"]
    income = facts.get("income") or {}
    balance = facts.get("balance") or {}
    market = facts.get("market") or {}
    ws = wb["历史财务数据"]
    last = FS_RIGHT_COL + len(hist)

    _orange_block(ws, 2, FS_LEFT_COL, ["利润表（亿元）", *hist])
    _orange_block(ws, 2, FS_RIGHT_COL, ["资产负债表（亿元）", *hist])

    left = 3
    for key in INCOME_ORDER:
        left = _write_fs_line(
            ws,
            left,
            FS_LEFT_COL,
            fs_label(key),
            hist,
            income.get(key),
            bold=key in INCOME_BOLD,
            formula=_income_formula(key, hist, ws) if key in {"毛利", "历史加权平均股本_百万股"} else None,
        )

    right = 3
    for key in BALANCE_ORDER:
        right = _write_fs_line(
            ws,
            right,
            FS_RIGHT_COL,
            fs_label(key),
            hist,
            balance.get(key),
            bold=key in BALANCE_BOLD,
            formula=_wc_formula(hist, ws) if key == "营运资本" else None,
        )

    metric_row = max(left, right) + 1
    _orange_block(ws, metric_row, FS_LEFT_COL, ["历史指标", *hist])
    row = metric_row + 1
    for i, (label, num_key, den_key) in enumerate(HIST_METRICS):
        ws.cell(row=row, column=FS_LEFT_COL, value=label)
        for j, year in enumerate(hist):
            col = FS_LEFT_COL + 1 + j
            if label == "营业收入增速":
                if j == 0:
                    continue
                prev = cell_ref(ws, fs_label(num_key), hist[j - 1], same_sheet=True)
                cur = cell_ref(ws, fs_label(num_key), year, same_sheet=True)
                ws.cell(row=row, column=col, value=f"={cur}/{prev}-1")
            else:
                num = cell_ref(ws, fs_label(num_key), year, same_sheet=True)
                den = cell_ref(ws, fs_label(den_key), year, same_sheet=True)
                ws.cell(row=row, column=col, value=f"={num}/{den}")
        row += 1

    row += 1
    _orange_block(ws, row, FS_LEFT_COL, ["市场快照"])
    row += 1
    row = _write_market(ws, row, market, facts.get("as_of"))

    write_sheet_title(ws, facts["company"], last)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["E"].width = 3
    ws.column_dimensions["F"].width = 28
    for col in range(2, 5):
        ws.column_dimensions[get_column_letter(col)].width = 12
    for col in range(7, last + 1):
        ws.column_dimensions[get_column_letter(col)].width = 12
    _format_hist(ws, last)


def _split(wb: Workbook, spec: dict) -> None:
    facts = spec["facts"]
    split = spec["split"]
    hist = facts["hist_periods"]
    ws = wb["主营业务拆分"]
    headers = ["分部名称", *hist, "单位", "备注"]
    last = len(headers)
    cells = [ws.cell(row=2, column=i, value=h) for i, h in enumerate(headers, 1)]
    apply_structure_style(ws, cells, "section")
    row = 3
    groups, company_res = build_split_groups(split)

    def write_leaf(seg: dict, *, indent: int = 0) -> None:
        nonlocal row
        cell = ws.cell(row=row, column=1, value=seg["name"])
        if indent:
            cell.alignment = Alignment(indent=indent, vertical="center")
        for i, val in enumerate(seg["hist_revenue"]):
            ws.cell(row=row, column=2 + i, value=val)
        ws.cell(row=row, column=2 + len(hist), value=seg["unit"])
        write_note(ws, row, last, seg["note"])
        row += 1

    for group in groups:
        if group.drilled:
            parent_row = row
            parent_cell = ws.cell(row=row, column=1, value=group.parent)
            parent_cell.font = Font(bold=True)
            unit = next((item.get("unit") for item in group.children if item.get("unit")), "亿元")
            ws.cell(row=row, column=2 + len(hist), value=unit)
            row += 1
            for child in group.children:
                write_leaf(child, indent=1)
            for i, year in enumerate(hist):
                parts = "+".join(
                    cell_ref(ws, child["name"], year, same_sheet=True) for child in group.children
                )
                ws.cell(row=parent_row, column=2 + i, value=f"={parts}")
        elif group.leaf:
            write_leaf(group.leaf)
    if company_res:
        write_leaf(company_res)
    ws.cell(row=row, column=1, value="合计")
    for i, year in enumerate(hist):
        parts = "+".join(
            cell_ref(ws, name, year, same_sheet=True) for name in top_level_names(groups, company_res)
        )
        ws.cell(row=row, column=2 + i, value=f"={parts}")
    row += 2
    write_section_row(ws, row, ["拆分说明"], last)
    row += 1
    for label, text in (
        ("拆分逻辑", split["split_logic"]),
        ("历史数据说明", split["hist_data_note"]),
        ("其他业务说明", split["other_note"]),
        ("主要来源", split["sources"]),
    ):
        ws.cell(row=row, column=1, value=label)
        write_note(ws, row, 2, text)
        row += 1
    write_sheet_title(ws, facts["company"], last)
    set_col_widths(ws, last, first=16, other=14)
    apply_number_formats(ws, set(), last, last)
    apply_number_formats(ws, set(), last, last)


def _revenue(wb: Workbook, spec: dict) -> None:
    facts = spec["facts"]
    split = spec["split"]
    forecasts = spec["forecasts"]
    hist = facts["hist_periods"]
    fcst = facts["forecast_periods"]
    periods = hist + fcst
    ws = wb["收入预测"]
    last = 2 + len(periods) + 1  # unit + 备注 → periods start at 2, unit = 2+len, 备注 = last
    write_first_section_header(ws, [*periods, "单位", "备注"])
    unit_col = last - 1
    note_col = last
    row = 3
    groups, company_res = build_split_groups(split)

    def write_assumption(seg: dict) -> None:
        nonlocal row
        name = seg["name"]
        method = seg["method"]
        fc = forecasts[name]
        spec = method_spec(method)
        unit_meta = fc.get("unit_meta") or {}
        write_group_row(ws, row, segment_group_title(name, method), last)
        row += 1
        rows: dict[str, int] = {}
        if spec.needs_fx:
            ws.cell(row=row, column=1, value="单位换算因子")
            for year in periods:
                ws.cell(row=row, column=_pcol(periods, year), value=unit_meta["fx"])
            ws.cell(row=row, column=unit_col, value=fx_unit_note(unit_meta))
            rows["fx"] = row
            row += 1
        for field in spec.assume:
            ws.cell(row=row, column=1, value=field)
            ws.cell(row=row, column=unit_col, value=assume_unit(field) or build_unit(field, unit_meta))
            if field == spec.assume_anchor:
                write_note(ws, row, note_col, fc.get("final_rationale") or "")
            rows[field] = row
            row += 1
        seg["_rows"] = rows

    for group in groups:
        if group.drilled:
            write_group_row(ws, row, group.parent, last)
            row += 1
            for child in group.children:
                write_assumption(child)
        elif group.leaf:
            write_assumption(group.leaf)
    if company_res:
        write_assumption(company_res)

    write_section_row(ws, row, ["收入构建区"], last)
    row += 1
    top_lines: list[tuple[str, int]] = []
    for group in groups:
        if group.drilled:
            write_group_row(ws, row, group.parent, last)
            row += 1
            parent_rev_row = row
            ws.cell(row=row, column=1, value="分部收入")
            ws.cell(row=row, column=unit_col, value="亿元")
            row += 1
            for child in group.children:
                row = _write_revenue_build(
                    ws, row, child, forecasts, hist, fcst, periods, last, unit_col
                )
            for year in periods:
                parts = "+".join(
                    cell_ref_at(ws, child["_rows"]["分部收入"], year)
                    for child in group.children
                )
                ws.cell(row=parent_rev_row, column=_pcol(periods, year), value=f"={parts}")
            top_lines.append((group.parent, parent_rev_row))
        elif group.leaf:
            row = _write_revenue_build(
                ws, row, group.leaf, forecasts, hist, fcst, periods, last, unit_col
            )
            top_lines.append((group.leaf["name"], group.leaf["_rows"]["分部收入"]))
    if company_res:
        row = _write_revenue_build(
            ws, row, company_res, forecasts, hist, fcst, periods, last, unit_col
        )
        top_lines.append((company_res["name"], company_res["_rows"]["分部收入"]))

    write_section_row(ws, row, ["收入归因"], last)
    row += 1
    ws.cell(row=row, column=1, value="营业收入")
    total_row = row
    for year in periods:
        parts = "+".join(cell_ref_at(ws, rev_row, year) for _, rev_row in top_lines)
        ws.cell(row=row, column=_pcol(periods, year), value=f"={parts}")
    row += 1
    for name, rev_row in top_lines:
        ws.cell(row=row, column=1, value=f"{name}收入占比")
        for year in periods:
            seg = cell_ref_at(ws, rev_row, year)
            total = cell_ref_at(ws, total_row, year)
            ws.cell(row=row, column=_pcol(periods, year), value=f"={seg}/{total}")
        row += 1

    write_sheet_title(ws, facts["company"], last)
    set_col_widths(ws, last, first=28, other=12)
    ws.column_dimensions[get_column_letter(note_col)].width = 42
    ratio = set()
    for r in range(3, ws.max_row + 1):
        lab = str(ws.cell(row=r, column=1).value or "")
        if any(token in lab for token in ("增速", "占比", "同比")) or lab.split("_", 1)[0] in RATE_FIELDS:
            ratio.add(r)
    apply_number_formats(ws, ratio, last, note_col)


def _write_revenue_build(ws, row, seg, forecasts, hist, fcst, periods, last, unit_col) -> int:
    name = seg["name"]
    method = seg["method"]
    fc = forecasts[name]
    spec = method_spec(method)
    unit_meta = fc.get("unit_meta") or {}
    rows = seg.setdefault("_rows", {})
    write_group_row(ws, row, segment_group_title(name, method), last)
    row += 1
    if spec.fcst_mode == "rev_grow":
        rev_row = row
        ws.cell(row=row, column=1, value="分部收入")
        for year in hist:
            ws.cell(row=row, column=_pcol(periods, year), value=fc["historical_data"]["分部收入"][year])
        ws.cell(row=row, column=unit_col, value="亿元")
        rows["分部收入"] = row
        row += 1
        grow_rev = rows["收入增速"]
        for i, year in enumerate(hist):
            if i == 0:
                _write_dash(ws, grow_rev, _pcol(periods, year))
                continue
            prev = hist[i - 1]
            cur = cell_ref_at(ws, rev_row, year)
            pri = cell_ref_at(ws, rev_row, prev)
            ws.cell(row=grow_rev, column=_pcol(periods, year), value=f"={cur}/{pri}-1")
        for i, year in enumerate(fcst):
            prev = hist[-1] if i == 0 else fcst[i - 1]
            g = cell_ref_at(ws, grow_rev, year)
            pri = cell_ref_at(ws, rev_row, prev)
            ws.cell(row=rev_row, column=_pcol(periods, year), value=f"={pri}*(1+{g})")
            ws.cell(
                row=grow_rev,
                column=_pcol(periods, year),
                value=as_growth_rate(fc["final_forecast"]["收入增速"][year]),
            )
    else:
        for field in spec.build:
            ws.cell(row=row, column=1, value=field)
            for year in hist:
                ws.cell(
                    row=row,
                    column=_pcol(periods, year),
                    value=fc["historical_data"][field][year],
                )
            ws.cell(row=row, column=unit_col, value=build_unit(field, unit_meta))
            rows[f"构建_{field}"] = row
            row += 1
        rev_row = row
        ws.cell(row=row, column=1, value="分部收入")
        rows["分部收入"] = row
        for year in hist + fcst:
            ws.cell(
                row=row,
                column=_pcol(periods, year),
                value=_driver_revenue_formula(ws, spec, rows, year),
            )
        ws.cell(row=row, column=unit_col, value="亿元")
        row += 1
        if spec.fcst_mode == "grow":
            for i, year in enumerate(fcst):
                prev = hist[-1] if i == 0 else fcst[i - 1]
                for field, grow in spec.grow_from.items():
                    prior = cell_ref_at(ws, rows[f"构建_{field}"], prev)
                    rate = cell_ref_at(ws, rows[grow], year)
                    ws.cell(
                        row=rows[f"构建_{field}"],
                        column=_pcol(periods, year),
                        value=f"={prior}*(1+{rate})",
                    )
            for field, grow in spec.grow_from.items():
                assume_row = rows[grow]
                for i, year in enumerate(hist):
                    if i == 0:
                        _write_dash(ws, assume_row, _pcol(periods, year))
                        continue
                    cur = cell_ref_at(ws, rows[f"构建_{field}"], year)
                    pri = cell_ref_at(ws, rows[f"构建_{field}"], hist[i - 1])
                    ws.cell(row=assume_row, column=_pcol(periods, year), value=f"={cur}/{pri}-1")
                for year in fcst:
                    ws.cell(
                        row=assume_row,
                        column=_pcol(periods, year),
                        value=as_growth_rate(fc["final_forecast"][grow][year]),
                    )
        else:
            for field in spec.build:
                assume_row = rows[field]
                build_row = rows[f"构建_{field}"]
                for year in hist:
                    src = cell_ref_at(ws, build_row, year)
                    ws.cell(row=assume_row, column=_pcol(periods, year), value=f"={src}")
                for year in fcst:
                    value = fc["final_forecast"][field][year]
                    if field in RATE_FIELDS:
                        value = as_growth_rate(value)
                    ws.cell(row=assume_row, column=_pcol(periods, year), value=value)
                    ws.cell(
                        row=build_row,
                        column=_pcol(periods, year),
                        value=f"={cell_ref_at(ws, assume_row, year)}",
                    )

    ws.cell(row=row, column=1, value="分部收入增速")
    rows["分部收入增速"] = row
    for i, year in enumerate(periods):
        if i == 0:
            _write_dash(ws, row, _pcol(periods, year))
            continue
        cur = cell_ref_at(ws, rows["分部收入"], year)
        pri = cell_ref_at(ws, rows["分部收入"], periods[i - 1])
        ws.cell(row=row, column=_pcol(periods, year), value=f"={cur}/{pri}-1")
    return row + 1


def _driver_revenue_formula(ws, spec, rows: dict[str, int], year: str) -> str:
    refs = {field: cell_ref_at(ws, rows[f"构建_{field}"], year) for field in spec.build}
    if spec.identity == "qty_price":
        fx = cell_ref_at(ws, rows["fx"], year)
        return f"={refs['销量']}*{refs['单价']}/{fx}"
    if spec.identity == "market":
        return f"={refs['市场规模']}*{refs['渗透率']}*{refs['市占率']}"
    if spec.identity == "order":
        return f"=({refs['期初在手订单']}+{refs['新签订单']})*{refs['转化率']}"
    if spec.identity == "store":
        fx = cell_ref_at(ws, rows["fx"], year)
        return f"={refs['门店数']}*{refs['单店产出']}/{fx}"
    if spec.identity == "user":
        fx = cell_ref_at(ws, rows["fx"], year)
        return f"={refs['订阅用户']}*{refs['单用户收入']}/{fx}"
    raise ValueError(f"没有收入公式: {spec.identity}")


def _cost(wb: Workbook, spec: dict) -> None:
    facts = spec["facts"]
    cost = spec["cost"]
    hist = facts["hist_periods"]
    fcst = facts["forecast_periods"]
    periods = hist + fcst
    ws = wb["运营成本预测"]
    last = 2 + len(periods)
    write_first_section_header(ws, [*periods, "备注"])
    row = 3

    def assume(label: str, hist_formula_from: str | None, fcst_map: dict, note: str = "") -> int:
        nonlocal row
        here = row
        ws.cell(row=row, column=1, value=label)
        if hist_formula_from:
            for year in hist:
                src = cell_ref(wb["历史财务数据"], hist_formula_from, year, same_sheet=False)
                if label == "有效税率":
                    den = cell_ref(wb["历史财务数据"], "税前利润", year, same_sheet=False)
                else:
                    den = cell_ref(wb["历史财务数据"], "营业收入", year, same_sheet=False)
                ws.cell(row=row, column=_pcol(periods, year), value=f"={src}/{den}")
        for year in fcst:
            ws.cell(row=row, column=_pcol(periods, year), value=fcst_map[year])
        if note:
            write_note(ws, row, last, note)
        row += 1
        return here

    assume("合并毛利率", "毛利", cost["gross_margin"], cost["rationale"])
    assume("销售费用率", "销售费用", cost["sales_ratio"])
    assume("管理费用率", "管理费用", cost["admin_ratio"])
    assume("研发费用率", "研发费用", cost["rd_ratio"])
    assume("有效税率", "所得税费用", cost["tax_rate"])
    ws.cell(row=row, column=1, value="财务费用假设")
    for year in fcst:
        ws.cell(row=row, column=_pcol(periods, year), value=cost["fin_exp"][year])
    row += 1
    ws.cell(row=row, column=1, value="营业外收入")
    for year in hist:
        ws.cell(
            row=row,
            column=_pcol(periods, year),
            value=f"={cell_ref(wb['历史财务数据'], '营业外收入', year)}",
        )
    for year in fcst:
        ws.cell(row=row, column=_pcol(periods, year), value=cost["nonop_inc"][year])
    row += 1
    ws.cell(row=row, column=1, value="营业外支出")
    for year in hist:
        ws.cell(
            row=row,
            column=_pcol(periods, year),
            value=f"={cell_ref(wb['历史财务数据'], '营业外支出', year)}",
        )
    for year in fcst:
        ws.cell(row=row, column=_pcol(periods, year), value=cost["nonop_exp"][year])
    row += 1
    ws.cell(row=row, column=1, value="少数股东比率")
    for year in fcst:
        ws.cell(row=row, column=_pcol(periods, year), value=cost["minority"][year])
    row += 1
    ws.cell(row=row, column=1, value="其他经营净收益")
    fs = wb["历史财务数据"]
    for year in hist:
        ebit = cell_ref(fs, "息税前利润", year)
        gp = cell_ref(fs, "毛利", year)
        s = cell_ref(fs, "销售费用", year)
        a = cell_ref(fs, "管理费用", year)
        d = cell_ref(fs, "研发费用", year)
        ws.cell(
            row=row,
            column=_pcol(periods, year),
            value=f"={ebit}-({gp}-{s}-{a}-{d})",
        )
    for year in fcst:
        ws.cell(row=row, column=_pcol(periods, year), value=cost["other_op"][year])
    row += 1

    write_section_row(ws, row, ["成本毛利区"], last)
    row += 1
    ws.cell(row=row, column=1, value="营业收入")
    for year in periods:
        ws.cell(row=row, column=_pcol(periods, year), value=f"={cell_ref(wb['收入预测'], '营业收入', year)}")
    row += 1
    ws.cell(row=row, column=1, value="毛利")
    for year in periods:
        rev = cell_ref(ws, "营业收入", year, same_sheet=True)
        gm = cell_ref(ws, "合并毛利率", year, same_sheet=True)
        ws.cell(row=row, column=_pcol(periods, year), value=f"={rev}*{gm}")
    row += 1
    ws.cell(row=row, column=1, value="营业成本")
    for year in periods:
        rev = cell_ref(ws, "营业收入", year, same_sheet=True)
        gp = cell_ref(ws, "毛利", year, same_sheet=True)
        ws.cell(row=row, column=_pcol(periods, year), value=f"={rev}-{gp}")
    row += 1

    write_section_row(ws, row, ["费用区"], last)
    row += 1
    for label, ratio in (
        ("销售费用", "销售费用率"),
        ("管理费用", "管理费用率"),
        ("研发费用", "研发费用率"),
    ):
        ws.cell(row=row, column=1, value=label)
        for year in periods:
            rev = cell_ref(ws, "营业收入", year, same_sheet=True)
            rt = cell_ref(ws, ratio, year, same_sheet=True)
            ws.cell(row=row, column=_pcol(periods, year), value=f"={rev}*{rt}")
        row += 1
    ws.cell(row=row, column=1, value="总费用")
    for year in periods:
        parts = "+".join(
            cell_ref(ws, fee, year, same_sheet=True)
            for fee in ("销售费用", "管理费用", "研发费用")
        )
        ws.cell(row=row, column=_pcol(periods, year), value=f"={parts}")

    write_sheet_title(ws, facts["company"], last)
    set_col_widths(ws, last, first=20)
    ws.column_dimensions[get_column_letter(last)].width = 36
    ratio_rows = set()
    for r in range(3, ws.max_row + 1):
        lab = str(ws.cell(row=r, column=1).value or "")
        if "率" in lab:
            ratio_rows.add(r)
    apply_number_formats(ws, ratio_rows, last, last)


def _pnl(wb: Workbook, spec: dict) -> None:
    facts = spec["facts"]
    hist = facts["hist_periods"]
    fcst = facts["forecast_periods"]
    periods = hist + fcst
    ws = wb["利润表预测"]
    last = 2 + len(periods)
    write_first_section_header(ws, [*periods, "备注"])
    labels = [
        "营业收入",
        "营业收入同比",
        "营业成本",
        "合并毛利率",
        "毛利",
        "销售费用",
        "管理费用",
        "研发费用",
        "息税前利润",
        "财务费用",
        "营业外收入",
        "营业外支出",
        "税前利润",
        "所得税费用",
        "少数股东损益",
        "归母净利润",
        "归母净利润同比",
        "EPS",
    ]
    derived = {"营业收入同比", "合并毛利率", "归母净利润同比"}
    row = 3
    for label in labels:
        ws.cell(row=row, column=1, value=label)
        if label not in derived:
            for year in hist:
                src = fs_label(label) if label == "EPS" else label
                ws.cell(
                    row=row,
                    column=_pcol(periods, year),
                    value=f"={cell_ref(wb['历史财务数据'], src, year)}",
                )
        row += 1

    def col(year: str) -> int:
        return _pcol(periods, year)

    for year in periods:
        ws.cell(
            row=_row(ws, "合并毛利率"),
            column=col(year),
            value=f"={cell_ref(wb['运营成本预测'], '合并毛利率', year)}",
        )
    for label, source in (
        ("营业收入同比", "营业收入"),
        ("归母净利润同比", "归母净利润"),
    ):
        r = _row(ws, label)
        for i, year in enumerate(periods):
            if i == 0:
                _write_dash(ws, r, col(year))
                continue
            cur = cell_ref(ws, source, year, same_sheet=True)
            pri = cell_ref(ws, source, periods[i - 1], same_sheet=True)
            ws.cell(row=r, column=col(year), value=f"={cur}/{pri}-1")

    for year in fcst:
        ws.cell(row=_row(ws, "营业收入"), column=col(year), value=f"={cell_ref(wb['收入预测'], '营业收入', year)}")
        ws.cell(row=_row(ws, "营业成本"), column=col(year), value=f"={cell_ref(wb['运营成本预测'], '营业成本', year)}")
        ws.cell(row=_row(ws, "毛利"), column=col(year), value=f"={cell_ref(wb['运营成本预测'], '毛利', year)}")
        for fee in ("销售费用", "管理费用", "研发费用"):
            ws.cell(row=_row(ws, fee), column=col(year), value=f"={cell_ref(wb['运营成本预测'], fee, year)}")
        gp = cell_ref(ws, "毛利", year, same_sheet=True)
        s = cell_ref(ws, "销售费用", year, same_sheet=True)
        a = cell_ref(ws, "管理费用", year, same_sheet=True)
        d = cell_ref(ws, "研发费用", year, same_sheet=True)
        oth = cell_ref(wb["运营成本预测"], "其他经营净收益", year)
        ws.cell(row=_row(ws, "息税前利润"), column=col(year), value=f"={gp}-{s}-{a}-{d}+{oth}")
        ws.cell(row=_row(ws, "财务费用"), column=col(year), value=f"={cell_ref(wb['运营成本预测'], '财务费用假设', year)}")
        ws.cell(row=_row(ws, "营业外收入"), column=col(year), value=f"={cell_ref(wb['运营成本预测'], '营业外收入', year)}")
        ws.cell(row=_row(ws, "营业外支出"), column=col(year), value=f"={cell_ref(wb['运营成本预测'], '营业外支出', year)}")
        ebit = cell_ref(ws, "息税前利润", year, same_sheet=True)
        fin = cell_ref(ws, "财务费用", year, same_sheet=True)
        ni = cell_ref(ws, "营业外收入", year, same_sheet=True)
        ne = cell_ref(ws, "营业外支出", year, same_sheet=True)
        ws.cell(row=_row(ws, "税前利润"), column=col(year), value=f"={ebit}-{fin}+{ni}-{ne}")
        pre = cell_ref(ws, "税前利润", year, same_sheet=True)
        taxr = cell_ref(wb["运营成本预测"], "有效税率", year)
        ws.cell(row=_row(ws, "所得税费用"), column=col(year), value=f"={pre}*{taxr}")
        tax = cell_ref(ws, "所得税费用", year, same_sheet=True)
        mino_ratio = cell_ref(wb["运营成本预测"], "少数股东比率", year)
        ws.cell(
            row=_row(ws, "少数股东损益"),
            column=col(year),
            value=f"=({pre}-{tax})*{mino_ratio}",
        )
        mino = cell_ref(ws, "少数股东损益", year, same_sheet=True)
        ws.cell(row=_row(ws, "归母净利润"), column=col(year), value=f"={pre}-{tax}-{mino}")
        ni2 = cell_ref(ws, "归母净利润", year, same_sheet=True)
        shares = f"'历史财务数据'!B{_row(wb['历史财务数据'], fs_label('当前总股本_百万股'))}"
        ws.cell(row=_row(ws, "EPS"), column=col(year), value=f"={ni2}*100/{shares}")

    write_note(
        ws,
        _row(ws, "税前利润"),
        last,
        "历史列沿用历史财务数据口径。预测列：息税前利润＝毛利－销售费用－管理费用－研发费用＋其他经营净收益；"
        "税前利润＝息税前利润－财务费用＋营业外收入－营业外支出。"
        "历史披露或接口的息税前利润口径可能与上述模型定义不同。",
    )
    write_sheet_title(ws, facts["company"], last)
    set_col_widths(ws, last, first=22)
    apply_number_formats(
        ws,
        {_row(ws, "营业收入同比"), _row(ws, "合并毛利率"), _row(ws, "归母净利润同比")},
        last,
        last,
    )


def _valuation(wb: Workbook, spec: dict) -> None:
    facts = spec["facts"]
    cons = spec["consensus"]
    comps = spec["comps"]
    fcst = facts["forecast_periods"]
    ws = wb["估值预测"]
    last = 6
    note_col = last
    write_first_section_header(ws, ["发布日期", *[f"{y}总营收" for y in fcst]])
    row = 3
    for inst in cons["institutions"]:
        ws.cell(row=row, column=1, value=inst["name"])
        ws.cell(row=row, column=2, value=inst["published"])
        ws.cell(row=row, column=2).number_format = "yyyy-mm-dd"
        for i, val in enumerate(inst["revenue"]):
            ws.cell(row=row, column=3 + i, value=val)
        row += 1
    median_row = row
    ws.cell(row=row, column=1, value="一致预期_中位数")
    first_inst = 3
    last_inst = row - 1
    for i in range(3):
        col = get_column_letter(3 + i)
        ws.cell(row=row, column=3 + i, value=consensus_median_formula(first_inst, last_inst, col))
    row += 1
    ws.cell(row=row, column=1, value="本模型预测")
    for i, year in enumerate(fcst):
        ws.cell(row=row, column=3 + i, value=f"={cell_ref(wb['收入预测'], '营业收入', year)}")
    row += 2

    y1 = fcst[0]
    write_section_row(ws, row, ["核心同业池", f"{y1} PE", "PE_TTM", "总市值", "单位", "备注"], last)
    row += 1
    core_start = row
    for name in comps["core"]:
        ws.cell(row=row, column=1, value=name["name"])
        ws.cell(row=row, column=2, value=name["pe_y1"])
        ws.cell(row=row, column=3, value=name["pe_ttm"])
        ws.cell(row=row, column=4, value=name["mcap"])
        ws.cell(row=row, column=5, value=_mcap_unit(name.get("currency")))
        write_note(ws, row, note_col, name.get("note") or "")
        row += 1
    core_end = row - 1
    ws.cell(row=row, column=1, value="核心池均值")
    ws.cell(row=row, column=2, value=f"=AVERAGE(B{core_start}:B{core_end})")
    ws.cell(row=row, column=3, value=f"=AVERAGE(C{core_start}:C{core_end})")
    row += 1
    ws.cell(row=row, column=1, value="核心池中位数")
    ws.cell(row=row, column=2, value=f"=MEDIAN(B{core_start}:B{core_end})")
    row += 1
    mean_pe = f"B{_row(ws, '核心池均值')}"
    eps = cell_ref(wb["利润表预测"], "EPS", y1)
    price = "历史财务数据!B" + str(_row(wb["历史财务数据"], fs_label("当前股价")))
    write_section_row(ws, row, ["综合估值预测", "PE调整系数", "目标PE", "目标价", "较现价空间"], last)
    row += 1
    neutral = row + 1
    scenario_rows = []
    for label, coef in (
        ("悲观", f"=B{neutral}*0.85"),
        ("中性", comps["pe_adjust"]),
        ("乐观", f"=B{neutral}*1.15"),
    ):
        scenario_rows.append(row)
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=coef)
        ws.cell(row=row, column=3, value=f"=B{row}*{mean_pe}")
        ws.cell(row=row, column=4, value=f"=C{row}*{eps}")
        ws.cell(row=row, column=5, value=f"=D{row}/{price}-1")
        row += 1
    ws.cell(row=row, column=1, value="目标PE")
    ws.cell(row=row, column=2, value=f"=C{neutral}")
    row += 1
    ws.cell(row=row, column=1, value="合理价值区间")
    bear, _neutral_row, bull = scenario_rows
    ws.cell(
        row=row,
        column=2,
        value=f'=TEXT(D{bear},"{NUM_FMT}")&" - "&TEXT(D{bull},"{NUM_FMT}")',
    )
    row += 1
    ws.cell(row=row, column=1, value="中枢目标价")
    ws.cell(row=row, column=2, value=f"=D{neutral}")
    row += 1
    ws.cell(row=row, column=1, value="较现价空间")
    space_row = row
    ws.cell(row=row, column=2, value=f"=E{neutral}")
    row += 1
    ws.cell(row=row, column=1, value="投资评级")
    ws.cell(row=row, column=2, value=f'=IF(B{space_row}>=0.15,"买入","观察")')

    _ = median_row
    write_sheet_title(ws, facts["company"], last)
    set_col_widths(ws, last, first=22, other=14)
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 36
    apply_number_formats(ws, {_row(ws, "较现价空间")}, last, None)
    for label in ("目标PE", "中枢目标价"):
        ws.cell(row=_row(ws, label), column=2).number_format = NUM_FMT
    range_cell = ws.cell(row=_row(ws, "合理价值区间"), column=2)
    range_cell.number_format = "@"
    range_cell.alignment = Alignment(horizontal="right")
    for scenario in scenario_rows:
        ws.cell(row=scenario, column=3).number_format = NUM_FMT
        ws.cell(row=scenario, column=4).number_format = NUM_FMT
        ws.cell(row=scenario, column=5).number_format = PCT_FMT


def _mcap_unit(currency: object) -> str:
    code = str(currency or "CNY").strip().upper() or "CNY"
    return {"CNY": "亿元", "HKD": "亿港元", "USD": "亿美元"}.get(code, code)


def _summary(wb: Workbook, spec: dict) -> None:
    from valuation.research_dossier.schema import DISCLAIMER, point_text, point_title

    facts = spec["facts"]
    notes = spec["summary_notes"]
    hist = facts["hist_periods"]
    fcst = facts["forecast_periods"]
    last_a = hist[-1]
    periods_show = [last_a, *fcst]
    periods_all = hist + fcst
    pnl = wb["利润表预测"]
    cost = wb["运营成本预测"]
    val = wb["估值预测"]
    fs = wb["历史财务数据"]
    ws = wb["总结"]
    last = 5
    row = 2
    row = _write_company_cover(ws, row, spec.get("company_info") or {}, last, fs)
    write_section_row(ws, row, ["核心结论"], last)
    row += 1
    for label in ("目标PE", "合理价值区间", "中枢目标价", "较现价空间", "投资评级"):
        ws.cell(row=row, column=1, value=label)
        src = _row(val, label)
        cell = ws.cell(row=row, column=2, value=f"=估值预测!B{src}")
        if label == "合理价值区间":
            cell.alignment = Alignment(horizontal="right")
        row += 1
    write_section_row(ws, row, ["关键指标"], last)
    row += 1
    ws.cell(row=row, column=1, value="当前股价")
    ws.cell(row=row, column=2, value=f"=历史财务数据!B{_row(fs, fs_label('当前股价'))}")
    row += 1
    ws.cell(row=row, column=1, value="EPS_最近实际")
    ws.cell(row=row, column=2, value=f"={cell_ref(pnl, 'EPS', last_a)}")
    row += 1
    ws.cell(row=row, column=1, value="EPS_预测首年")
    ws.cell(row=row, column=2, value=f"={cell_ref(pnl, 'EPS', fcst[0])}")
    row += 1
    write_section_row(ws, row, ["财务预测概览", *periods_show], last)
    row += 1
    for label in ("营业收入", "毛利", "息税前利润", "归母净利润", "EPS"):
        ws.cell(row=row, column=1, value=label)
        for i, year in enumerate(periods_show):
            ws.cell(row=row, column=2 + i, value=f"={cell_ref(pnl, label, year)}")
        row += 1
    ws.cell(row=row, column=1, value="营业收入同比增速")
    for i, year in enumerate(periods_show):
        prev = periods_all[periods_all.index(year) - 1]
        cur = cell_ref(pnl, "营业收入", year)
        pri = cell_ref(pnl, "营业收入", prev)
        ws.cell(row=row, column=2 + i, value=f"={cur}/{pri}-1")
    row += 1
    ws.cell(row=row, column=1, value="息税前利润率")
    for i, year in enumerate(periods_show):
        ebit = cell_ref(pnl, "息税前利润", year)
        rev = cell_ref(pnl, "营业收入", year)
        ws.cell(row=row, column=2 + i, value=f"={ebit}/{rev}")
    row += 1
    ws.cell(row=row, column=1, value="归母净利润率")
    for i, year in enumerate(periods_show):
        ni = cell_ref(pnl, "归母净利润", year)
        rev = cell_ref(pnl, "营业收入", year)
        ws.cell(row=row, column=2 + i, value=f"={ni}/{rev}")
    row += 1
    write_section_row(ws, row, ["关键经营假设", *periods_show], last)
    row += 1
    for label in ("合并毛利率", "销售费用率", "管理费用率", "研发费用率", "有效税率"):
        ws.cell(row=row, column=1, value=label)
        for i, year in enumerate(periods_show):
            ws.cell(row=row, column=2 + i, value=f"={cell_ref(cost, label, year)}")
        row += 1
    for title, key in (
        ("投研逻辑", "thesis"),
        ("未来展望", "outlook"),
        ("主要风险", "risks"),
    ):
        write_section_row(ws, row, [title], last)
        row += 1
        for i, point in enumerate(notes.get(key) or [], 1):
            title = point_title(point)
            ws.cell(row=row, column=1, value=f"{i}. {title}".strip())
            write_note(ws, row, 2, point_text(point))
            ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last)
            ws.row_dimensions[row].height = 36
            row += 1
    cell = ws.cell(row=row, column=1, value=DISCLAIMER)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    apply_structure_style(ws, [cell], "footnote")
    cell.font = Font(color="888888", italic=True)
    cell.alignment = Alignment(wrap_text=True, vertical="center")

    write_sheet_title(ws, facts["company"], last)
    set_col_widths(ws, last, first=22, other=14)
    ws.column_dimensions["B"].width = 36
    ratio_rows = {
        _row(ws, "较现价空间"),
        _row(ws, "营业收入同比增速"),
        _row(ws, "息税前利润率"),
        _row(ws, "归母净利润率"),
        _row(ws, "合并毛利率"),
        _row(ws, "销售费用率"),
        _row(ws, "管理费用率"),
        _row(ws, "研发费用率"),
        _row(ws, "有效税率"),
    }
    apply_number_formats(ws, ratio_rows, last, None)
    for label in ("投资评级", "合理价值区间"):
        ws.cell(row=_row(ws, label), column=2).number_format = "@"


def _write_company_cover(ws, row: int, info: dict, last: int, fs) -> int:
    profile = company_profile_rows(info)
    market = company_market_rows(info)
    if not profile and not market:
        return row
    write_section_row(ws, row, ["公司信息"], last)
    row += 1
    wrap_labels = {"主营业务"}
    for label, value in profile:
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=value)
        if label in wrap_labels:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last)
            ws.row_dimensions[row].height = 48
        row += 1
    if market:
        write_section_row(ws, row, ["当前行情"], last)
        row += 1
        for label, value, is_pct in market:
            ws.cell(row=row, column=1, value=label)
            if label == "当前股价（元）":
                ws.cell(row=row, column=2, value=f"=历史财务数据!B{_row(fs, fs_label('当前股价'))}")
            elif label == "总市值（亿元）":
                ws.cell(row=row, column=2, value=f"=历史财务数据!B{_row(fs, fs_label('总市值_亿'))}")
            elif label == "流通市值（亿元）":
                ws.cell(row=row, column=2, value=f"=历史财务数据!B{_row(fs, fs_label('流通市值_亿'))}")
            else:
                ws.cell(row=row, column=2, value=value)
            if is_pct:
                ws.cell(row=row, column=2).number_format = PCT_FMT
            elif label in {"当前股价（元）", "总市值（亿元）", "流通市值（亿元）"}:
                ws.cell(row=row, column=2).number_format = NUM_FMT
            row += 1
    return row


def _orange_block(ws, row: int, start: int, values: list) -> None:
    cells = [
        ws.cell(row=row, column=start + i, value=value)
        for i, value in enumerate(values)
    ]
    apply_structure_style(ws, cells, "section")
    for cell in cells:
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _bold_block(ws, row: int, start: int, width: int) -> None:
    for col in range(start, start + width):
        ws.cell(row=row, column=col).font = Font(bold=True)


def _write_fs_line(
    ws,
    row: int,
    label_col: int,
    label: str,
    hist: list[str],
    series,
    *,
    bold: bool = False,
    formula: list[str] | None = None,
) -> int:
    ws.cell(row=row, column=label_col, value=label)
    for i, _year in enumerate(hist):
        col = label_col + 1 + i
        if formula and i < len(formula) and formula[i] is not None:
            ws.cell(row=row, column=col, value=formula[i])
        elif series and i < len(series) and series[i] is not None:
            ws.cell(row=row, column=col, value=series[i])
    if bold:
        _bold_block(ws, row, label_col, 1 + len(hist))
    return row + 1


def _income_formula(key: str, hist: list[str], ws) -> list[str] | None:
    if key == "毛利":
        try:
            find_row_by_label(ws, "营业收入")
            find_row_by_label(ws, "营业成本")
        except ValueError:
            return None
        return [
            f"={cell_ref(ws, '营业收入', year, same_sheet=True)}-{cell_ref(ws, '营业成本', year, same_sheet=True)}"
            for year in hist
        ]
    if key == "历史加权平均股本_百万股":
        ni = fs_label("归母净利润")
        eps = fs_label("EPS")
        return [
            f"={cell_ref(ws, ni, year, same_sheet=True)}*100/{cell_ref(ws, eps, year, same_sheet=True)}"
            for year in hist
        ]
    return None


def _wc_formula(hist: list[str], ws) -> list[str] | None:
    try:
        find_row_by_label(ws, "流动资产合计")
        find_row_by_label(ws, "流动负债合计")
    except ValueError:
        return None
    return [
        f"={cell_ref(ws, '流动资产合计', year, same_sheet=True)}-{cell_ref(ws, '流动负债合计', year, same_sheet=True)}"
        for year in hist
    ]


def _write_market(ws, row: int, market: dict, as_of) -> int:
    written: dict[str, int] = {}
    for key in MARKET_ROWS:
        label = fs_label(key)
        if key == "总市值_亿":
            price_row = written.get("当前股价")
            share_row = written.get("当前总股本_百万股")
            if price_row and share_row:
                ws.cell(row=row, column=1, value=label)
                ws.cell(row=row, column=2, value=f"=B{price_row}*B{share_row}/100")
                written[key] = row
                row += 1
            continue
        if key == "流通市值_亿":
            price_row = written.get("当前股价")
            float_row = written.get("流通股本_百万股")
            if price_row and float_row:
                ws.cell(row=row, column=1, value=label)
                ws.cell(row=row, column=2, value=f"=B{price_row}*B{float_row}/100")
                written[key] = row
                row += 1
            continue
        value = market.get(key)
        if key == "快照日期" and not value:
            value = str(as_of)[:10] if as_of else None
        if value is None or value == "":
            continue
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        written[key] = row
        row += 1
    return row


def _format_hist(ws, last: int) -> None:
    for r in range(3, ws.max_row + 1):
        for label_col in (FS_LEFT_COL, FS_RIGHT_COL):
            label = str(ws.cell(row=r, column=label_col).value or "")
            if not label:
                continue
            for col in range(label_col + 1, min(label_col + 4, last + 1)):
                cell = ws.cell(row=r, column=col)
                if cell.value is None:
                    continue
                if isinstance(cell.value, str) and not str(cell.value).startswith("="):
                    continue
                if any(token in label for token in ("率", "增速")):
                    cell.number_format = PCT_FMT
                else:
                    cell.number_format = NUM_FMT


def _col(ws, header: str) -> int:
    return find_col_by_header(ws, header)


def _write_dash(ws, row: int, column: int) -> None:
    cell = ws.cell(row=row, column=column, value="-")
    cell.data_type = "s"
    cell.alignment = Alignment(horizontal="right")


def _row(ws, label: str) -> int:
    return find_row_by_label(ws, label)


def _pcol(periods: list[str], year: str) -> int:
    return 2 + periods.index(year)


def consensus_median_formula(first_inst: int, last_inst: int, col: str) -> str | None:
    if last_inst < first_inst:
        return None
    return f"=MEDIAN({col}{first_inst}:{col}{last_inst})"
