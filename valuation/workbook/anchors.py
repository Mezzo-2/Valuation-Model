"""
openpyxl 工作簿脚本用的 anchor 工具。

用途：
- 用稳定行标签解析行 anchor，避免硬编码行号。
- 用表头文本解析列/年份 anchor。
- 构建可靠的跨表公式链接。
"""

from __future__ import annotations

from dataclasses import dataclass
import re
from typing import Iterable, Optional

from openpyxl.cell.cell import Cell
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string


_PERIOD_HEADER_RE = re.compile(r"^\d{4}(?:A|E|Q[1-4][AE])$")

# 历史财务数据左右分栏：利润表/指标在 A，资产负债表在 F。
_FS_LABEL_COLS = (1, 6)
_FS_YEAR_SPAN = 3


@dataclass(frozen=True)
class CellAnchor:
    sheet: str
    row: int
    col: int

    @property
    def a1(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"


def _normalize_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().lower()


def _quote_sheet_name(name: str) -> str:
    """
    为 Excel 公式引用对工作表名加引号。

    Excel 规则：用单引号包裹表名；表名内的单引号须加倍转义
    （例如 ``Q4'24`` -> ``'Q4''24'``）。
    始终加引号可避免空格、前导数字等边界情况。
    """
    return "'" + name.replace("'", "''") + "'"


def _fs_label_cols(ws: Worksheet, label_col: int) -> tuple[int, ...]:
    if ws.title == "历史财务数据" and label_col == 1:
        return _FS_LABEL_COLS
    return (label_col,)


def find_label_position(
    ws: Worksheet,
    label: str,
    label_col: int = 1,
    min_row: int = 1,
    max_row: Optional[int] = None,
    exact: bool = True,
) -> tuple[int, int]:
    """返回 (row, label_col)。历史财务数据会扫 A 列和 F 列。"""
    target = _normalize_text(label)
    scan_min = min_row
    scan_max = ws.max_row if max_row is None else max_row
    matches: list[tuple[int, int]] = []
    for col in _fs_label_cols(ws, label_col):
        for r in range(scan_min, scan_max + 1):
            value = _normalize_text(ws.cell(row=r, column=col).value)
            if not value:
                continue
            if (exact and value == target) or (not exact and target in value):
                matches.append((r, col))
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise ValueError(
            f"Duplicate label in sheet '{ws.title}': '{label}' "
            f"at {matches} (label_col={label_col})"
        )
    raise ValueError(
        f"Label not found in sheet '{ws.title}': '{label}' (label_col={label_col})"
    )


def find_row_by_label(
    ws: Worksheet,
    label: str,
    label_col: int = 1,
    min_row: int = 1,
    max_row: Optional[int] = None,
    exact: bool = True,
) -> int:
    """
    在标签列（默认 A 列）按文本匹配查找行号。

    Args:
        ws: openpyxl 工作表
        label: 目标标签文本
        label_col: 标签列索引（1 起）
        min_row: 扫描下限（含）
        max_row: 扫描上限（默认 ws.max_row）；``None`` 表示用 ws.max_row，
            ``0`` 视为无行可扫并抛出 ValueError
        exact: True 为精确匹配，False 为子串匹配
    """
    row, _ = find_label_position(
        ws,
        label,
        label_col=label_col,
        min_row=min_row,
        max_row=max_row,
        exact=exact,
    )
    return row


def find_header_row(
    ws: Worksheet,
    required_header: Optional[str] = None,
    *,
    min_row: int = 1,
    max_scan_rows: int = 50,
) -> int:
    """
    自动定位列头行，不假定第 1 行是表头。

    - 传入 ``required_header`` 时，返回唯一包含该表头的行；缺失或多行命中即报错。
    - 未传入时，寻找至少包含两个期间列头（如 ``2025A``、``2026E``）的唯一行。

    A1 可安全保留为工作表标题；期间型工作表通常会识别到第 2 行。
    非期间宽表应传入一个稳定的 ``required_header``。
    """
    if min_row < 1 or max_scan_rows < 1:
        raise ValueError("min_row and max_scan_rows must be positive integers")

    scan_max = min(ws.max_row, min_row + max_scan_rows - 1)
    matches: list[int] = []
    target = _normalize_text(required_header)
    for r in range(min_row, scan_max + 1):
        values = [
            _normalize_text(ws.cell(row=r, column=c).value)
            for c in range(1, ws.max_column + 1)
        ]
        if required_header is not None:
            if target in values:
                matches.append(r)
        else:
            period_count = sum(
                1 for value in values if _PERIOD_HEADER_RE.fullmatch(value.upper())
            )
            if period_count >= 2:
                matches.append(r)

    criterion = (
        f"required header '{required_header}'"
        if required_header is not None
        else "at least two period headers"
    )
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise ValueError(
            f"Ambiguous header row in sheet '{ws.title}': {criterion} "
            f"matched rows {matches}"
        )
    raise ValueError(
        f"Header row not found in sheet '{ws.title}': {criterion} "
        f"within rows {min_row}..{scan_max}"
    )


def find_col_by_header(
    ws: Worksheet,
    header: str,
    header_row: Optional[int] = None,
    min_col: int = 1,
    max_col: Optional[int] = None,
    exact: bool = True,
) -> int:
    """
    在指定表头行按文本匹配查找列号。

    ``max_col=None`` 使用 ws.max_column；``max_col=0`` 抛出 ValueError。
    """
    if header_row is None:
        header_row = find_header_row(ws, required_header=header)
    target = _normalize_text(header)
    scan_max = ws.max_column if max_col is None else max_col
    matches: list[int] = []
    for c in range(min_col, scan_max + 1):
        value = _normalize_text(ws.cell(row=header_row, column=c).value)
        if not value:
            continue
        if (exact and value == target) or (not exact and target in value):
            matches.append(c)
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise ValueError(
            f"Duplicate header in sheet '{ws.title}': '{header}' at columns "
            f"{matches} (header_row={header_row})"
        )
    raise ValueError(
        f"Header not found in sheet '{ws.title}': '{header}' (header_row={header_row})"
    )


def _header_row_in_block(
    ws: Worksheet,
    label_row: int,
    label_col: int,
    header: str,
) -> int:
    target = _normalize_text(header)
    min_c = label_col + 1
    max_c = label_col + _FS_YEAR_SPAN
    for r in range(label_row, 0, -1):
        for c in range(min_c, max_c + 1):
            if _normalize_text(ws.cell(row=r, column=c).value) == target:
                return r
    raise ValueError(
        f"Header '{header}' not found near label row {label_row} "
        f"in sheet '{ws.title}'"
    )


def get_anchor(
    ws: Worksheet,
    label: str,
    header: str,
    *,
    label_col: int = 1,
    header_row: Optional[int] = None,
    exact: bool = True,
) -> CellAnchor:
    """
    通过（行标签 + 列表头）解析单元格 anchor。
    """
    row, lc = find_label_position(ws, label, label_col=label_col, exact=exact)
    if ws.title == "历史财务数据":
        hr = header_row or _header_row_in_block(ws, row, lc, header)
        col = find_col_by_header(
            ws,
            header,
            header_row=hr,
            min_col=lc + 1,
            max_col=lc + _FS_YEAR_SPAN,
            exact=exact,
        )
    else:
        col = find_col_by_header(ws, header, header_row=header_row, exact=exact)
    return CellAnchor(sheet=ws.title, row=row, col=col)


def build_link(anchor: CellAnchor, lock_col: bool = False, lock_row: bool = False) -> str:
    """
    由 anchor 生成跨表公式链接字符串。

    表名始终加单引号，内嵌单引号加倍转义，以支持 ``Q4'24``、``Cost & Margin`` 等名称。

    输出示例：
      ='收入构建'!F42
      ='收入构建'!$F$42
      ='Q4''24'!F42
    """
    col_letter = get_column_letter(anchor.col)
    if lock_col:
        col_letter = f"${col_letter}"
    row_ref = f"${anchor.row}" if lock_row else str(anchor.row)
    return f"={_quote_sheet_name(anchor.sheet)}!{col_letter}{row_ref}"


def write_link(
    target_ws: Worksheet,
    target_row: int,
    target_col: int,
    source_anchor: CellAnchor,
    *,
    lock_col: bool = False,
    lock_row: bool = False,
) -> None:
    """
    在目标单元格写入跨表链接公式。
    """
    target_ws.cell(row=target_row, column=target_col).value = (
        build_link(source_anchor, lock_col=lock_col, lock_row=lock_row)
    )


def find_last_full_year_col(
    ws: Worksheet,
    header_row: Optional[int] = None,
    min_col: int = 2,
    max_col: Optional[int] = None,
) -> int:
    """
    定位最后一个完整年度实际数据列（后缀 A 且不含 Q）。

    当期间 schema 包含季度列（如 2026Q1A）时，forecast 公式的基期必须引用
    最后一个完整年度列，而非物理上最后一个 A 列。

    匹配规则：列头匹配 ``\\d{4}A``（如 2025A）但不匹配 ``Q\\d``。

    Raises:
        ValueError: 未找到任何完整年度列。
    """
    if header_row is None:
        header_row = find_header_row(ws)
    full_year_re = re.compile(r"^\d{4}A$")
    scan_max = ws.max_column if max_col is None else max_col
    last_col: Optional[int] = None
    for c in range(min_col, scan_max + 1):
        hdr = str(ws.cell(row=header_row, column=c).value or "").strip()
        if full_year_re.match(hdr):
            last_col = c
    if last_col is None:
        raise ValueError(
            f"No full-year actual column found in sheet '{ws.title}' "
            f"(header_row={header_row}, scanned cols {min_col}..{scan_max})"
        )
    return last_col


def find_required_rows(
    ws: Worksheet,
    labels: Iterable[str],
    *,
    label_col: int = 1,
    exact: bool = True,
) -> dict[str, int]:
    """
    一次解析多个行标签。
    任一标签缺失即抛出 ValueError。
    """
    return {
        label: find_row_by_label(ws, label, label_col=label_col, exact=exact)
        for label in labels
    }


# ─────────────────────────────────────────────────────────────
# 反公式错位（anti-misalignment）
# 核心原则：公式里的每个引用都必须由「标签 + 表头」当场解析，
# 禁止任何行号算术（row ± k）、禁止记忆整数、禁止硬编码 A1。
# ─────────────────────────────────────────────────────────────

def cell_ref(
    ws: Worksheet,
    label: str,
    header: str,
    *,
    label_col: int = 1,
    header_row: Optional[int] = None,
    lock_col: bool = False,
    lock_row: bool = False,
    exact: bool = True,
    same_sheet: bool = False,
) -> str:
    """
    由「行标签 + 列表头」即时解析出单元格引用字符串，供拼接进公式。

    这是写公式的唯一推荐方式：永远不要手拼列字母+行号，也不要用
    find_row_by_label(...) ± k 这类偏移。

    Args:
        same_sheet: True 时返回不带表名的引用（如 ``F4``），用于同表内引用；
            False 时返回带表名的跨表引用（如 ``'运营成本'!F4``）。

    Returns:
        引用字符串（不含前导 ``=``），如 ``'收入构建'!F25`` 或 ``$F$4``。
    """
    row, lc = find_label_position(ws, label, label_col=label_col, exact=exact)
    if ws.title == "历史财务数据":
        hr = header_row or _header_row_in_block(ws, row, lc, header)
        col = find_col_by_header(
            ws,
            header,
            header_row=hr,
            min_col=lc + 1,
            max_col=lc + _FS_YEAR_SPAN,
            exact=exact,
        )
    else:
        col = find_col_by_header(ws, header, header_row=header_row, exact=exact)
    col_letter = get_column_letter(col)
    if lock_col:
        col_letter = f"${col_letter}"
    row_ref = f"${row}" if lock_row else str(row)
    a1 = f"{col_letter}{row_ref}"
    if same_sheet:
        return a1
    return f"{_quote_sheet_name(ws.title)}!{a1}"


def cell_ref_at(
    ws: Worksheet,
    row: int,
    header: str,
    *,
    same_sheet: bool = True,
    header_row: Optional[int] = None,
) -> str:
    """已知行号时，只按列表头解析列，拼出 A1。"""
    col = find_col_by_header(ws, header, header_row=header_row)
    a1 = f"{get_column_letter(col)}{row}"
    if same_sheet:
        return a1
    return f"{_quote_sheet_name(ws.title)}!{a1}"


_SECTION_TITLES = frozenset({"收入驱动假设区", "收入构建区", "收入归因"})


def find_row_in_group(
    ws: Worksheet,
    group_title: str,
    label: str,
    *,
    after_row: int = 1,
    until_row: Optional[int] = None,
    group_titles: Optional[set[str]] = None,
    label_col: int = 1,
) -> int:
    """在指定分组标题下查找短标签，不跨下一分组或下一节。"""
    scan_max = ws.max_row if until_row is None else until_row - 1
    target = _normalize_text(group_title)
    header_row: Optional[int] = None
    for r in range(after_row, scan_max + 1):
        if _normalize_text(ws.cell(row=r, column=label_col).value) == target:
            header_row = r
            break
    if header_row is None:
        raise ValueError(
            f"Group '{group_title}' not found in sheet '{ws.title}' "
            f"within rows {after_row}..{scan_max}"
        )
    titles = set(group_titles or ())
    end = scan_max
    for r in range(header_row + 1, scan_max + 1):
        val = str(ws.cell(row=r, column=label_col).value or "").strip()
        if val in _SECTION_TITLES:
            end = r - 1
            break
        if val != group_title and (
            val in titles or ("（" in val and val.endswith("）"))
        ):
            end = r - 1
            break
    return find_row_by_label(
        ws, label, label_col=label_col, min_row=header_row + 1, max_row=end
    )


def group_title_above(
    ws: Worksheet,
    row: int,
    group_titles: set[str],
    label_col: int = 1,
) -> str:
    """从该行往上找最近的分组标题。"""
    for r in range(row - 1, 0, -1):
        val = str(ws.cell(row=r, column=label_col).value or "").strip()
        if val in group_titles:
            return val
        if val in _SECTION_TITLES:
            return ""
    return ""


def get_val(
    wb: Workbook,
    sheet: str,
    label: str,
    header: str,
    *,
    label_col: int = 1,
    header_row: Optional[int] = None,
    exact: bool = True,
) -> object:
    """
    按工作表、canonical 行标签和列表头读取单元格值。

    工作表、标签或表头不存在时抛出 ``KeyError`` / ``ValueError``，
    不返回模糊的 ``None`` 作为“未找到”信号。单元格本身为空时仍返回 ``None``。
    """
    ws = wb[sheet]
    anchor = get_anchor(
        ws,
        label,
        header,
        label_col=label_col,
        header_row=header_row,
        exact=exact,
    )
    return ws.cell(row=anchor.row, column=anchor.col).value


def write_note(
    ws: Worksheet,
    row: int,
    notes_col: int,
    text: object,
) -> None:
    """
    将备注安全写为纯文本，避免 Excel 将说明解析为公式。

    通过强制 ``data_type="s"`` 让单元格始终以字符串类型落盘：即使文本以
    ``=`` / ``+`` / ``-`` / ``@`` 开头，Excel/WPS 也按文本显示，不会解析成公式
    产生 ``#NAME?``。**刻意不加前导单引号**——单引号会被 openpyxl 原样写入，
    在单元格里显示成多余的 ``'``，且会让 ``备注`` 值不再以 ``=`` 开头，从而绕过
    Rule 36 验证门（掩盖真正应改写为自然语言的问题备注）。
    """
    note = "" if text is None else str(text).strip()
    cell = ws.cell(row=row, column=notes_col)
    cell.value = note
    cell.data_type = "s"


# 备注列禁止出现的「工程化语言」子串（§3.9）：规则/步骤编号、JSON 字段、条件逻辑箭头。
_NOTE_ENG_SUBSTRINGS = (
    "final_rationale",
    "final_forecast",
    "forecast_notes",
    "->",
    "→",
)
# 规则/步骤编号（Rule 32 / Step 6，大小写与空格不敏感）。
_NOTE_RULE_STEP_RE = re.compile(r"(?:rule|step)\s*\d", re.IGNORECASE)
# sheet 路径引用：`估值预测!D8`、`历史财务数据!营业收入` 等（`!` 后接列字母/`$`/中文/字母）。
_NOTE_SHEET_PATH_RE = re.compile(r"\S+!\s*[\$A-Za-z\u4e00-\u9fff]")
# 其余 snake_case JSON 字段（如 some_field），至少两段全小写字母。
_NOTE_SNAKE_RE = re.compile(r"\b[a-z]{3,}_[a-z]{3,}\b")


def note_is_engineering(text: object) -> bool:
    """
    判断一段备注是否含「工程化语言」——供 Rule 36 验证门扫描 ``备注`` 列使用。

    命中任一即视为不合规（应改写为自然语言）：
    - 以 ``=`` / ``+`` / ``-`` / ``@`` 开头，或文本任意位置含 ``=``（公式式表达）；
    - Excel 错误字面量（``#NAME?`` / ``#VALUE!`` / ``#REF!`` / ``#DIV/0!`` / ``#NUM!``）；
    - sheet 路径引用（``sheet!Cell`` / ``sheet!标签``）；
    - 规则/步骤编号（``Rule 32`` / ``Step 6``）；
    - JSON 字段名或 snake_case 标识（``final_rationale`` 等）；
    - 条件逻辑箭头（``->`` / ``→``）。

    质量标签（``[有据可查]`` 等中文方括号标签）不受影响，可正常出现在开头。
    """
    t = "" if text is None else str(text).strip()
    if not t:
        return False
    if t[0] in "=+-@":
        return True
    if "=" in t:
        return True
    error_literals = {"#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#NUM!", "#NULL!"}
    if t in error_literals or any(e in t for e in error_literals):
        return True
    low = t.lower()
    if any(s in low for s in _NOTE_ENG_SUBSTRINGS):
        return True
    if _NOTE_RULE_STEP_RE.search(t):
        return True
    if _NOTE_SHEET_PATH_RE.search(t):
        return True
    if _NOTE_SNAKE_RE.search(low):
        return True
    return False


def verify_formula_refs(
    ws: Worksheet,
    formula_cell: Cell,
    expected_label_row_map: dict[str, int],
) -> None:
    """
    验证公式引用的行号包含所有预期 canonical 行。

    ``expected_label_row_map`` 的值应由 ``find_row_by_label`` 当场解析，
    不得传入记忆或硬编码的行号。非公式单元格直接抛出 ``ValueError``。
    """
    formula = str(formula_cell.value or "")
    if not formula.startswith("="):
        raise ValueError(
            f"Expected formula at {ws.title}!{formula_cell.coordinate}, got {formula!r}"
        )
    row_nums = {
        int(match)
        for match in re.findall(r"\$?[A-Z]{1,3}\$?(\d+)", formula)
    }
    for label, expected_row in expected_label_row_map.items():
        if expected_row not in row_nums:
            raise ValueError(
                f"FORMULA REF ERROR at {ws.title}!{formula_cell.coordinate}: "
                f"expected row {expected_row} ('{label}') not found in "
                f"formula '{formula}'. Found rows: {sorted(row_nums)}"
            )


_FORMULA_REF_RE = re.compile(
    r"(?:(?:'((?:[^']|'')+)'|([A-Za-z0-9_\u4e00-\u9fff]+))!)?"
    r"\$?([A-Z]{1,3})\$?(\d+)"
)


def verify_formula_anchors(
    ws: Worksheet,
    formula_cell: Cell,
    expected_anchors: Iterable[CellAnchor],
) -> None:
    """验证公式同时引用了预期 sheet、行与列；跨表公式优先使用此函数。"""
    formula = str(formula_cell.value or "")
    if not formula.startswith("="):
        raise ValueError(
            f"Expected formula at {ws.title}!{formula_cell.coordinate}, got {formula!r}"
        )

    found: set[tuple[str, int, int]] = set()
    previous_match_end = -1
    previous_sheet = ws.title
    for match in _FORMULA_REF_RE.finditer(formula):
        quoted_sheet, bare_sheet, col_letters, row_text = match.groups()
        explicit_sheet = quoted_sheet.replace("''", "'") if quoted_sheet else bare_sheet
        between = formula[previous_match_end : match.start()] if previous_match_end >= 0 else ""
        # Excel 只在区间起点写一次表名：'Sheet'!A1:B2。冒号后的 B2
        # 继承起点表名；其他未限定引用仍属于公式所在工作表。
        sheet = previous_sheet if explicit_sheet is None and between.strip() == ":" else (
            explicit_sheet or ws.title
        )
        found.add((sheet, int(row_text), column_index_from_string(col_letters)))
        previous_sheet = sheet
        previous_match_end = match.end()

    for anchor in expected_anchors:
        key = (anchor.sheet, anchor.row, anchor.col)
        if key not in found:
            rendered = sorted(f"{sheet}!{get_column_letter(col)}{row}" for sheet, row, col in found)
            raise ValueError(
                f"FORMULA ANCHOR ERROR at {ws.title}!{formula_cell.coordinate}: "
                f"expected {anchor.sheet}!{anchor.a1} not found in {formula!r}; "
                f"found {rendered}"
            )
