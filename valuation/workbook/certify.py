from __future__ import annotations

import math
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.worksheet import Worksheet

from valuation.historical_financials.subjects import fs_label
from valuation.segment_research.methods import formula_errors, method_spec
from valuation.segment_split.schema import t_recon
from valuation.workbook.anchors import (
    CellAnchor,
    find_col_by_header,
    find_row_by_label,
    find_row_in_group,
    group_title_above,
    verify_formula_anchors,
)
from valuation.workbook.build import consensus_median_formula
from valuation.workbook.build import SHEET_ORDER
from valuation.workbook.load_json import load_spec
from valuation.workbook.tree import (
    build_split_groups,
    revenue_group_titles,
    segment_group_title,
)

_A1_RE = re.compile(r"(?:'[^']+'!)?(\$?[A-Z]{1,3}\$?\d+)")
_DANGER_FORMULA = re.compile(r"#DIV/0!|#REF!|/\s*0\b", re.I)
_PNL_CALC_LABELS = ("毛利", "息税前利润", "归母净利润", "EPS")


def formula_ref_rows(ws: Worksheet, formula: str) -> list[tuple[int, str]]:
    """把公式里的 A1 引用还原成 (行号, A 列标签)。"""
    refs: list[tuple[int, str]] = []
    for token in _A1_RE.findall(str(formula).replace(" ", "")):
        _col, row = coordinate_from_string(token.replace("$", ""))
        refs.append((row, str(ws.cell(row=row, column=1).value or "")))
    return refs


def formula_row_labels(ws: Worksheet, formula: str) -> list[str]:
    """把公式里的 A1 引用还原成该行 A 列标签。"""
    return [label for _row, label in formula_ref_rows(ws, formula)]


def certify_workbook(run_dir: Path, workbook_path: Path) -> dict:
    spec = load_spec(run_dir)
    wb = load_workbook(workbook_path, data_only=False)
    calc_fails: list[str] = []
    format_fails: list[str] = []
    fails = calc_fails
    titles = [ws.title for ws in wb.worksheets]
    if titles != SHEET_ORDER:
        format_fails.append(f"tab 顺序 {titles} != {SHEET_ORDER}")

    for name in SHEET_ORDER:
        if name not in wb.sheetnames:
            format_fails.append(f"缺 sheet {name}")

    rev = wb["收入预测"]
    try:
        r = find_row_by_label(rev, "营业收入")
        sample = rev.cell(row=r, column=5).value
        if not (isinstance(sample, str) and sample.startswith("=")):
            fails.append("收入预测!营业收入 预测列不是公式")
    except ValueError as exc:
        fails.append(str(exc))

    group_titles = revenue_group_titles(spec["split"])
    try:
        build_start = find_row_by_label(rev, "收入构建区") + 1
        build_end = find_row_by_label(rev, "收入归因")
    except ValueError as exc:
        fails.append(str(exc))
        build_start, build_end = 1, None

    for seg in spec["split"]["final_segments"]:
        title = segment_group_title(seg["name"], seg["method"])
        label = f"分部收入/{seg['name']}"
        try:
            r = find_row_in_group(
                rev,
                title,
                "分部收入",
                after_row=build_start,
                until_row=build_end,
                group_titles=group_titles,
            )
        except ValueError as exc:
            fails.append(str(exc))
            continue
        formula = str(rev.cell(row=r, column=5).value or "")
        for item in formula_errors(seg["method"], formula):
            fails.append(f"{label} {item}")
        if method_spec(seg["method"]).needs_fx:
            labels = formula_row_labels(rev, formula)
            if not any(lab == "单位换算因子" or lab.startswith("单位换算因子") for lab in labels):
                fails.append(f"{label} 分部收入公式未引用单位换算因子行: {formula}")

    groups, _company_res = build_split_groups(spec["split"])
    for group in groups:
        if not group.drilled:
            continue
        label = f"分部收入/{group.parent}"
        try:
            r = find_row_in_group(
                rev,
                group.parent,
                "分部收入",
                after_row=build_start,
                until_row=build_end,
                group_titles=group_titles,
            )
        except ValueError as exc:
            fails.append(str(exc))
            continue
        formula = str(rev.cell(row=r, column=5).value or "")
        if "(1+" in formula.replace(" ", ""):
            fails.append(f"{label} 应是子项加总，不能是增速外推: {formula}")
        ref_groups = {
            group_title_above(rev, row, group_titles)
            for row, lab in formula_ref_rows(rev, formula)
            if lab == "分部收入"
        }
        for child in group.children:
            child_title = segment_group_title(child["name"], child["method"])
            if child_title not in ref_groups:
                fails.append(f"{label} 未引用 {child['name']}")
    try:
        r = find_row_by_label(rev, "营业收入")
        formula = str(rev.cell(row=r, column=5).value or "")
        ref_groups = {
            group_title_above(rev, row, group_titles)
            for row, lab in formula_ref_rows(rev, formula)
            if lab == "分部收入"
        }
        for group in groups:
            if not group.drilled:
                continue
            for child in group.children:
                child_title = segment_group_title(child["name"], child["method"])
                if child_title in ref_groups:
                    fails.append(f"营业收入 不应引用子项 {child['name']}")
    except ValueError as exc:
        fails.append(str(exc))

    pnl = wb["利润表预测"]
    y1 = spec["facts"]["forecast_periods"][0]
    try:
        r = find_row_by_label(pnl, "少数股东损益")
        c = find_col_by_header(pnl, y1)
        cell = pnl.cell(row=r, column=c).value
        text = str(cell or "")
        if not (isinstance(cell, str) and text.startswith("=")):
            fails.append(f"利润表预测!少数股东损益 {y1} 不是公式: {cell!r}")
        elif "少数股东比率" not in text and "运营成本预测" not in text:
            fails.append(f"利润表预测!少数股东损益 {y1} 未引用少数股东比率: {text}")
    except ValueError as exc:
        fails.append(str(exc))

    val = wb["估值预测"]
    for label in ("目标PE", "目标价", "上行空间", "投资评级"):
        try:
            r = find_row_by_label(val, label)
            cell = val.cell(row=r, column=2).value
            if not (isinstance(cell, str) and str(cell).startswith("=")):
                fails.append(f"估值预测!{label} 不是公式")
        except ValueError as exc:
            fails.append(str(exc))

    sm = wb["总结"]
    formula_labels = (
        "目标PE",
        "目标价",
        "上行空间",
        "投资评级",
        "当前股价",
        "EPS_最近实际",
        "EPS_预测首年",
        "营业收入",
        "营业收入同比增速",
        "毛利",
        "息税前利润",
        "息税前利润率",
        "归母净利润",
        "EPS",
        "合并毛利率",
        "销售费用率",
        "管理费用率",
        "研发费用率",
        "有效税率",
    )
    for label in formula_labels:
        try:
            r = find_row_by_label(sm, label)
            cell = sm.cell(row=r, column=2).value
            if not (isinstance(cell, str) and str(cell).startswith("=")):
                fails.append(f"总结!{label} 不是公式")
        except ValueError as exc:
            fails.append(str(exc))

    narrative = ("公司背景", "投研逻辑", "未来展望", "主要风险")
    for title in narrative:
        try:
            find_row_by_label(sm, title)
        except ValueError as exc:
            format_fails.append(str(exc))
    if any(
        "数据缺口" in str(sm.cell(row=r, column=c).value or "")
        for r in range(1, (sm.max_row or 1) + 1)
        for c in range(1, 6)
    ):
        format_fails.append("总结出现数据缺口")
    last_text = str(sm.cell(row=sm.max_row, column=1).value or "")
    if "不构成投资建议" not in last_text:
        format_fails.append("总结最后一行不是免责声明")

    fs = wb["历史财务数据"]
    if fs.cell(row=2, column=1).value != "利润表（亿元）":
        format_fails.append(f"历史财务数据!A2={fs.cell(row=2, column=1).value}")
    if fs.cell(row=2, column=6).value != "资产负债表（亿元）":
        format_fails.append(f"历史财务数据!F2={fs.cell(row=2, column=6).value}")
    if fs.max_column > 4 and any(
        str(fs.cell(row=2, column=c).value or "") == "备注" for c in range(1, fs.max_column + 1)
    ):
        format_fails.append("历史财务数据仍有备注列表头")

    calc_fails.extend(hist_recon_errors(spec["facts"], spec["split"]["final_segments"]))
    calc_fails.extend(_pnl_year_errors(wb, spec))
    calc_fails.extend(_consensus_median_errors(wb, spec))
    calc_fails.extend(_snapshot_replay_errors(spec))

    pes = [float(c["pe_y1"]) for c in spec["comps"]["core"]]
    report = {
        "pass": not calc_fails,
        "fails": calc_fails + format_fails,
        "calc_fails": calc_fails,
        "format_fails": format_fails,
        "core_pe_mean": sum(pes) / len(pes) if pes else None,
        "workbook": str(workbook_path),
    }
    return report


def formula_is_dangerous(text: str) -> bool:
    return bool(_DANGER_FORMULA.search(str(text or "")))


def median_range_is_empty(formula: str) -> bool:
    match = re.search(
        r"MEDIAN\(([A-Z]+)(\d+):\1(\d+)\)",
        str(formula or "").replace("$", "").replace(" ", ""),
        re.I,
    )
    if not match:
        return False
    return int(match.group(3)) < int(match.group(2))


def _pnl_year_errors(wb, spec: dict) -> list[str]:
    fails: list[str] = []
    pnl = wb["利润表预测"]
    fcst = list(spec["facts"]["forecast_periods"])
    y1 = fcst[0]
    for year in fcst:
        try:
            col = find_col_by_header(pnl, year)
        except ValueError as exc:
            fails.append(str(exc))
            continue
        for label in _PNL_CALC_LABELS:
            try:
                row = find_row_by_label(pnl, label)
            except ValueError as exc:
                fails.append(str(exc))
                continue
            cell = pnl.cell(row=row, column=col)
            text = str(cell.value or "")
            if not (isinstance(cell.value, str) and text.startswith("=")):
                fails.append(f"利润表预测!{label} {year} 不是公式: {cell.value!r}")
                continue
            if formula_is_dangerous(text):
                fails.append(f"利润表预测!{label} {year} 公式危险: {text}")
    try:
        eps_row = find_row_by_label(pnl, "EPS")
        ni_row = find_row_by_label(pnl, "归母净利润")
        col = find_col_by_header(pnl, y1)
        shares_row = find_row_by_label(wb["历史财务数据"], fs_label("当前总股本_百万股"))
        verify_formula_anchors(
            pnl,
            pnl.cell(row=eps_row, column=col),
            [
                CellAnchor("利润表预测", ni_row, col),
                CellAnchor("历史财务数据", shares_row, 2),
            ],
        )
    except ValueError as exc:
        fails.append(str(exc))
    return fails


def _consensus_median_errors(wb, spec: dict) -> list[str]:
    fails: list[str] = []
    val = wb["估值预测"]
    try:
        row = find_row_by_label(val, "一致预期_中位数")
    except ValueError as exc:
        return [str(exc)]
    institutions = (spec.get("consensus") or {}).get("institutions") or []
    for i in range(3):
        cell = val.cell(row=row, column=3 + i)
        text = str(cell.value or "")
        if not institutions:
            if text.upper().startswith("=MEDIAN") or median_range_is_empty(text):
                fails.append(f"估值预测一致预期中位数在无机构时不应写 MEDIAN: {text}")
            continue
        if median_range_is_empty(text):
            fails.append(f"估值预测一致预期中位数区间为空: {text}")
        expected = consensus_median_formula(3, 2 + len(institutions), chr(ord("C") + i))
        _ = expected
    return fails


def _snapshot_replay_errors(spec: dict) -> list[str]:
    try:
        from valuation.research_dossier.snapshot import snapshot_from_spec

        snap = snapshot_from_spec(spec)
    except Exception as exc:
        return [f"snapshot 回放失败: {exc}"]
    y1 = spec["facts"]["forecast_periods"][0]
    fails: list[str] = []
    rev = (snap.get("revenue") or {}).get(y1)
    eps = ((snap.get("pnl") or {}).get(y1) or {}).get("eps")
    if rev is None or not math.isfinite(float(rev)):
        fails.append(f"snapshot {y1} 营收不是有限数")
    if eps is None or not math.isfinite(float(eps)):
        fails.append(f"snapshot {y1} EPS 不是有限数")
    return fails


def hist_recon_errors(facts: dict, segs: list[dict]) -> list[str]:
    fails: list[str] = []
    inc = facts["income"]
    for i, year in enumerate(facts["hist_periods"]):
        pre = inc["税前利润"][i]
        chk = inc["营业利润"][i] + inc["营业外收入"][i] - inc["营业外支出"][i]
        if abs(pre - chk) > t_recon(float(pre)):
            fails.append(f"{year} 税前利润勾稽失败 {pre} vs {chk}")
        total = inc["营业收入"][i]
        summed = sum(float(seg["hist_revenue"][i]) for seg in segs)
        if abs(summed - float(total)) > t_recon(float(total)):
            fails.append(f"{year} 分部加总 {summed} != {total}")
    return fails
