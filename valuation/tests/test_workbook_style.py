from __future__ import annotations

import unittest

from openpyxl import Workbook

from valuation.workbook.anchors import cell_ref, find_row_by_label
from valuation.workbook.build import _cost
from valuation.workbook.styles import (
    CROSS_GREEN,
    FORMULA_BLACK,
    GROUP,
    HEADER,
    INK,
    INPUT_BLUE,
    NEAR,
    ON_SECTION,
    PAGE,
    PINGFANG,
    SECTION,
    TEXT_FONT,
    WHITE,
    YAHEI,
    choose_number_format,
    finish_sheet,
    formula_role_color,
    mark_input,
    reset_style_marks,
    write_block_header,
    write_column_header_row,
    write_first_section_header,
    write_group_row,
    write_page_title,
)


def _rgb(color) -> str:
    if color is None:
        return ""
    rgb = getattr(color, "rgb", None)
    if not isinstance(rgb, str):
        return ""
    return rgb[-6:].upper()


class NumberFormatTests(unittest.TestCase):
    def test_choose_number_format(self):
        cases = {
            "营业收入": '#,##0.00;[Red](#,##0.00);0.00;@',
            "合并毛利率": "0.00%;[Red](0.00%);0.00%;@",
            "营业收入同比增速": "0.00%;[Red](0.00%);0.00%;@",
            "较现价空间": "+0.00%;[Red]-0.00%;0.00%;@",
            "1日涨跌幅": "+0.00%;[Red]-0.00%;0.00%;@",
            "目标PE": '0.00"x";[Red](0.00"x");0.00"x";@',
            "PE_TTM": '0.00"x";[Red](0.00"x");0.00"x";@',
            "PB": '0.00"x";[Red](0.00"x");0.00"x";@',
            "EPS": "0.00;[Red](0.00);0.00;@",
            "EPS_预测首年": "0.00;[Red](0.00);0.00;@",
            "当前股价": "0.00;[Red](0.00);0.00;@",
            "中枢目标价": "0.00;[Red](0.00);0.00;@",
            "当前总股本（百万股）": '#,##0.00;[Red](#,##0.00);0.00;@',
            "单位换算因子": '#,##0.00;[Red](#,##0.00);0.00;@',
            "合理价值区间": "@",
            "投资评级": "@",
        }
        for label, fmt in cases.items():
            self.assertEqual(choose_number_format(label), fmt, label)


class HeaderStyleTests(unittest.TestCase):
    def test_hist_and_forecast_headers_differ(self):
        wb = Workbook()
        ws = wb.active
        write_column_header_row(ws, 3, ["", "2025A", "2026E", "2027E"])
        self.assertEqual(_rgb(ws["B3"].fill.fgColor), NEAR)
        self.assertEqual(_rgb(ws["C3"].fill.fgColor), HEADER)
        self.assertEqual(_rgb(ws["D3"].fill.fgColor), HEADER)
        self.assertTrue(ws["C3"].border.left is None or ws["C3"].border.left.style is None)
        self.assertTrue(ws["D3"].border.left is None or ws["D3"].border.left.style is None)

    def test_module_title_sits_above_year_headers(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "利润表预测"
        nxt = write_first_section_header(ws, ["2025A", "2026E"])
        self.assertEqual(nxt, 3)
        self.assertEqual(ws["A2"].value, "利润表预测")
        self.assertEqual(ws["B2"].value, "2025A")
        self.assertEqual(ws["C2"].value, "2026E")
        self.assertEqual(_rgb(ws["A2"].fill.fgColor), SECTION)
        self.assertEqual(_rgb(ws["B2"].fill.fgColor), SECTION)
        self.assertEqual(_rgb(ws["C2"].fill.fgColor), SECTION)
        self.assertEqual(_rgb(ws["A2"].font.color), WHITE)
        self.assertFalse(ws["A2"].alignment.indent)

    def test_group_row_is_not_limited_to_revenue_sheet(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "总结"
        write_group_row(ws, 3, "当前行情", 4)
        self.assertEqual(ws["A3"].value, "当前行情")
        self.assertEqual(_rgb(ws["A3"].fill.fgColor), GROUP)
        self.assertEqual(_rgb(ws["A3"].font.color), ON_SECTION)

    def test_page_title_is_white_on_deep_green(self):
        reset_style_marks()
        wb = Workbook()
        ws = wb.active
        write_page_title(ws, "英维克（002837）", 4)
        finish_sheet(ws)
        self.assertEqual(ws["A1"].value, "英维克（002837）")
        self.assertEqual(ws["A1"].alignment.horizontal, "center")
        self.assertEqual(ws["A1"].alignment.vertical, "center")
        self.assertEqual(_rgb(ws["A1"].font.color), WHITE)
        self.assertEqual(_rgb(ws["A1"].fill.fgColor), PAGE)
        self.assertEqual(_rgb(ws["D1"].fill.fgColor), PAGE)
        self.assertTrue(ws["A1"].border.bottom is None or ws["A1"].border.bottom.style is None)
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertIn(TEXT_FONT, {YAHEI, PINGFANG})


class ValueRoleTests(unittest.TestCase):
    def test_formula_role_color(self):
        self.assertEqual(formula_role_color("=B2*2"), FORMULA_BLACK)
        self.assertEqual(formula_role_color("=利润表预测!B4"), CROSS_GREEN)
        self.assertEqual(formula_role_color("=[其他簿.xlsx]利润表!B4"), "FF0000")

    def test_inputs_are_blue_and_history_is_not(self):
        reset_style_marks()
        wb = Workbook()
        ws = wb.active
        ws.title = "运营成本预测"
        ws["A3"] = "合并毛利率"
        ws["B3"] = 0.25
        ws["C3"] = 0.27
        mark_input(ws, 3, 3)
        ws["D3"] = "=B3"
        ws["E3"] = "=历史财务数据!B8"
        finish_sheet(ws)
        self.assertEqual(_rgb(ws["B3"].font.color), FORMULA_BLACK)
        self.assertEqual(_rgb(ws["C3"].font.color), INPUT_BLUE)
        self.assertEqual(_rgb(ws["D3"].font.color), FORMULA_BLACK)
        self.assertEqual(_rgb(ws["E3"].font.color), CROSS_GREEN)
        self.assertEqual(_rgb(ws["A3"].font.color), INK)
        self.assertFalse(ws.sheet_view.showGridLines)


class CostHistoryFormulaTests(unittest.TestCase):
    def test_fin_exp_and_minority_history_link_to_statements(self):
        reset_style_marks()
        wb = Workbook()
        fs = wb.active
        fs.title = "历史财务数据"
        write_block_header(fs, 2, 1, ["利润表（亿元）", "2025A"])
        for index, label in enumerate(
            (
                "营业收入",
                "毛利",
                "销售费用",
                "管理费用",
                "研发费用",
                "财务费用",
                "营业外收入",
                "营业外支出",
                "税前利润",
                "所得税费用",
                "归母净利润",
                "少数股东损益",
                "息税前利润",
            )
        ):
            fs.cell(row=3 + index, column=1, value=label)
            fs.cell(row=3 + index, column=2, value=1)
        revenue = wb.create_sheet("收入预测")
        revenue["A2"] = "收入驱动假设区"
        revenue["B2"] = "2025A"
        revenue["C2"] = "2026E"
        revenue["A3"] = "营业收入"
        wb.create_sheet("运营成本预测")
        _cost(
            wb,
            {
                "facts": {
                    "company": "测试",
                    "hist_periods": ["2025A"],
                    "forecast_periods": ["2026E"],
                },
                "cost": {
                    "gross_margin": {"2026E": 0.1},
                    "sales_ratio": {"2026E": 0.01},
                    "admin_ratio": {"2026E": 0.01},
                    "rd_ratio": {"2026E": 0.01},
                    "tax_rate": {"2026E": 0.1},
                    "fin_exp": {"2026E": 1.0},
                    "nonop_inc": {"2026E": 0.0},
                    "nonop_exp": {"2026E": 0.0},
                    "minority": {"2026E": -0.02},
                    "other_op": {"2026E": 0.0},
                    "rationale": "2026E毛利率0.1。",
                },
            },
        )
        ws = wb["运营成本预测"]
        fin = find_row_by_label(ws, "财务费用假设")
        self.assertEqual(ws.cell(row=fin, column=2).value, f"={cell_ref(fs, '财务费用', '2025A')}")
        self.assertEqual(ws.cell(row=fin, column=3).value, 1.0)
        minority_row = find_row_by_label(ws, "少数股东比率")
        minority = cell_ref(fs, "少数股东损益", "2025A")
        parent = cell_ref(fs, "归母净利润", "2025A")
        self.assertEqual(
            ws.cell(row=minority_row, column=2).value,
            f"={minority}/({parent}+{minority})",
        )
        self.assertEqual(ws.cell(row=minority_row, column=3).value, -0.02)


if __name__ == "__main__":
    unittest.main()
