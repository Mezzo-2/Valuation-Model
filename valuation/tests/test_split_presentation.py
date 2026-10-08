from __future__ import annotations

import unittest

from openpyxl import Workbook

from valuation.segment_split.hist_agent import (
    HistCell, HistFill, HistFillDeps, HistSeg, SourceDocument, hist_errors, notes_from_draft,
)
from valuation.segment_split.presentation import source_files
from valuation.workbook.adapt import adapt_split
from valuation.workbook.build import _split


class SplitPresentationTests(unittest.TestCase):
    def setUp(self):
        self.facts = {
            "company": "测试公司", "ticker": "000001", "unit": "亿元",
            "hist_periods": ["2024A", "2025A"],
            "income": {"营业收入": [100.0, 110.0]},
        }
        self.plan = {
            "company": "测试公司", "ticker": "000001", "caliber": "official",
            "official_parents": ["硬件", "软件"],
            "segments": [
                {"name": "硬件", "parent": "硬件"},
                {"name": "软件", "parent": "软件"},
                {"name": "其他", "parent": ""},
            ],
        }
        self.official = {"item_classify": "按产品", "rows": [
            {"period": "2025A", "name": "硬件", "publish_date": "2026-03-31"},
            {"period": "2025A", "name": "软件", "publish_date": "2026-03-31"},
            {"period": "2024A", "name": "硬件", "publish_date": "2025-03-28"},
        ]}
        self.draft = HistFill(
            fill_method="官方抄录",
            segments=[
                HistSeg(name="硬件", years=[
                    HistCell(year="2024A", amount=60, quality="直接披露", why="年报披露数"),
                    HistCell(year="2025A", amount=70, quality="直接披露", why="年报披露数"),
                ]),
                HistSeg(name="软件", years=[
                    HistCell(year="2024A", amount=30, quality="有据可查", why="券商报告"),
                    HistCell(year="2025A", amount=25, quality="直接披露", why="年报披露数"),
                ]),
                HistSeg(name="其他", years=[
                    HistCell(year="2024A", amount=10, quality="倒推", why="收入减已列分部"),
                    HistCell(year="2025A", amount=15, quality="倒推", why="收入减已列分部"),
                ]),
            ],
            split_logic="按年报产品口径保留硬件、软件；未单列金额列入其他。",
            other_note="其他包含未单列产品，具体构成未知。",
            source_documents=[SourceDocument(title="业务跟踪", date="2025-08-01", institution="甲证券")],
        )

    def test_rebuilds_explanations_from_verified_numbers_and_sources(self):
        notes = notes_from_draft(
            self.facts, self.plan, self.official, self.draft,
            hits=[{"title": "业务跟踪", "date": "2025-08-01", "institution": "甲证券"}],
        )
        split = adapt_split(notes, self.facts)
        remarks = {seg["name"]: seg["note"] for seg in split["final_segments"]}
        self.assertEqual(remarks["硬件"], "年报披露")
        self.assertEqual(remarks["软件"], "2024A券商资料；2025A年报披露")
        self.assertEqual(remarks["其他"], "倒推")
        self.assertIn("硬件：占比63.6%，同比+16.7%，对收入变动贡献+10.00亿元", split["structure_change"])
        self.assertIn("软件：占比22.7%，同比-16.7%", split["structure_change"])
        self.assertIn("公司营业收入与已列分部之间的差额", split["other_note"])
        self.assertIn("2026-03-31", split["source_files"])
        self.assertIn("甲证券－业务跟踪｜2025-08-01", split["source_files"])
        self.assertNotIn("年报披露数；年报披露数", str(remarks))

        wb = Workbook()
        wb.active.title = "主营业务拆分"
        _split(wb, {"facts": self.facts, "split": split})
        ws = wb.active
        self.assertEqual([ws.cell(2, i).value for i in range(1, 6)], ["分部名称", "2024A", "2025A", "单位", "备注"])
        labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        self.assertEqual([x for x in labels if x in {"拆分逻辑", "结构变化", "其他业务说明", "来源文件"}],
                         ["拆分逻辑", "结构变化", "其他业务说明", "来源文件"])
        self.assertNotIn("历史数据说明", labels)
        self.assertNotIn("主要来源", labels)
        row = labels.index("结构变化") + 1
        self.assertTrue(ws.cell(row, 2).alignment.wrap_text)
        for title in ("拆分逻辑", "结构变化", "其他业务说明", "来源文件"):
            title_row = labels.index(title) + 1
            title_cell = ws.cell(title_row, 1)
            detail_cell = ws.cell(title_row, 2)
            self.assertTrue(title_cell.font.bold)
            self.assertEqual(title_cell.alignment.horizontal, "left")
            self.assertEqual(title_cell.alignment.vertical, "center")
            self.assertEqual(detail_cell.alignment.horizontal, "left")
            self.assertEqual(detail_cell.alignment.vertical, "center")
            self.assertTrue(any(r.min_col == 2 and r.max_col == 5 and r.min_row == title_row
                                for r in ws.merged_cells.ranges))
        structure_row = labels.index("结构变化") + 1
        self.assertTrue(any(r.min_col == 1 and r.max_col == 1 and r.min_row == structure_row
                            and r.max_row > structure_row for r in ws.merged_cells.ranges))
        self.assertTrue(any(r.height > 18 for r in ws.row_dimensions.values() if r.height))

    def test_unverified_document_is_not_listed(self):
        notes = notes_from_draft(self.facts, self.plan, self.official, self.draft, hits=[])
        self.assertNotIn("业务跟踪", source_files(notes))

    def test_drilled_split_requires_a_verified_document(self):
        deps = HistFillDeps(
            company="测试公司", ticker="000001", hist_periods=["2025A"],
            forecast_years=["2026"], revenue=[110.0],
            plan_names=["硬件", "软件", "其他"], plan_segments=self.plan["segments"],
            official_names=[], official_rows=[], search=lambda *_args: [],
            plan_caliber="drilled",
            prior_hits=[{"title": "业务跟踪", "date": "2025-08-01"}],
        )
        no_docs = self.draft.model_copy(update={"source_documents": []})
        self.assertTrue(any("须列至少一份" in error for error in hist_errors(deps, no_docs)))
        self.assertEqual(hist_errors(deps, self.draft), [])


if __name__ == "__main__":
    unittest.main()
