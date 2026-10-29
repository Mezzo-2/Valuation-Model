"""港美默认沿用 A 股的快照股本和具体收支，避免重复假设或静默丢失调整。"""
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from valuation.historical_financials.foreign import decode_financials
from valuation.operating_cost.cost_agent import CostDraft
from valuation.operating_cost.schema import require_cost
from valuation.segment_split.schema import SchemaError
from valuation.shared.markets import foreign_pnl, register_labels, snapshot_ordinary_shares
from valuation.tests.test_foreign_financials import facts, cost, integration_spec, replies
from valuation.workbook.anchors import find_row_by_label
from valuation.workbook.build import build_workbook
from valuation.workbook.certify import certify_workbook


class MarketSimplificationTests(unittest.TestCase):
    def test_default_workbook_has_no_duplicate_share_or_profit_rows(self):
        forbidden = {"经营利润净调整", "营业利润至税前净调整", "普通股归属扣减", "普通股股东利润",
                     "预测基本加权平均普通股数（百万股）", "其他经营净收益", "优先股股息"}
        for market in ["HK", "US"]:
            with self.subTest(market=market), tempfile.TemporaryDirectory() as directory:
                spec = integration_spec(market)
                with patch("valuation.workbook.build.load_spec", return_value=spec):
                    path = build_workbook(Path(directory))
                wb = load_workbook(path)
                register_labels(wb, spec["facts"])
                for name in ["运营成本预测", "利润表预测", "总结"]:
                    labels = {str(c.value) for row in wb[name] for c in row}
                    self.assertFalse(forbidden & labels, (name, forbidden & labels))
                pnl, fs = wb["利润表预测"], wb["历史财务数据"]
                eps = pnl.cell(find_row_by_label(pnl, "EPS"), 5).value
                self.assertIn(f"'历史财务数据'!B{find_row_by_label(fs,'当前总股本_百万股')}", eps)
                self.assertIn("营业外收入", {row[0].value for row in pnl})

    def test_snapshot_share_change_updates_eps_and_ads_ratio_is_used_once(self):
        f = facts("US")
        c = cost(f)
        revenue = {y: 100 for y in f["hist_periods"]+f["forecast_periods"]}
        before = foreign_pnl(f, c, revenue)["2026E"]["eps"]
        f["market"]["每交易单位普通股数"] = 8
        self.assertEqual(snapshot_ordinary_shares(f), 8000)
        self.assertAlmostEqual(foreign_pnl(f,c,revenue)["2026E"]["eps"],before/8)
        f["market"]["当前总股本_百万股"] *= 2
        self.assertAlmostEqual(foreign_pnl(f,c,revenue)["2026E"]["eps"],before/16)

    def test_ads_workbook_references_ratio_in_financial_snapshot(self):
        spec = integration_spec("US")
        spec["facts"]["market"]["每交易单位普通股数"] = 8
        with tempfile.TemporaryDirectory() as directory:
            with patch("valuation.workbook.build.load_spec", return_value=spec):
                path = build_workbook(Path(directory))
            with patch("valuation.workbook.certify.load_spec", return_value=spec):
                report = certify_workbook(Path(directory), path)
            self.assertEqual(report["fails"], [])
            wb = load_workbook(path)
            register_labels(wb, spec["facts"])
            fs,pnl = wb["历史财务数据"],wb["利润表预测"]
            ratio = f"'历史财务数据'!B{find_row_by_label(fs,'每交易单位普通股数')}"
            self.assertEqual(pnl.cell(find_row_by_label(pnl,"EPS"),5).value.count(ratio),1)

    def test_actual_profit_items_keep_expense_and_income_signs(self):
        f = facts("HK")
        c = cost(f)
        c["fin_exp"] = {y:-2 for y in f["forecast_periods"]}
        c["associates"] = {y: .5 for y in f["forecast_periods"]}
        c = require_cost(c,f)
        result = foreign_pnl(f,c,{y:100 for y in f["hist_periods"]+f["forecast_periods"]})["2026E"]
        self.assertEqual(result["pre"]-result["ebit"],3.5)
        del c["associates"]["2027E"]
        with self.assertRaises(SchemaError):
            require_cost(c,f)

    def test_missing_applicable_income_is_not_silently_zero(self):
        f = facts("HK")
        c = cost(f)
        c.pop("nonop_inc")
        with self.assertRaisesRegex(SchemaError,"财务及其他非经营收支"):
            require_cost(c,f)
        f["income"]["财务费用"] = [None,1,2]
        c["nonop_inc"] = {y:1 for y in f["forecast_periods"]}
        with self.assertRaisesRegex(SchemaError,"fin_exp"):
            require_cost(c,f)

    def test_nonzero_legacy_adjustments_or_changed_shares_cannot_be_dropped(self):
        f = facts("HK")
        for key,value in [("pretax_adjustment",1),("common_adjustment",1),("forecast_shares",2000)]:
            with self.subTest(key=key):
                c = cost(f)
                c[key] = {y:value for y in f["forecast_periods"]}
                with self.assertRaises(SchemaError):
                    require_cost(c,f)
        self.assertNotIn("forecast_shares", CostDraft.model_fields)
        self.assertNotIn("common_adjustment", CostDraft.model_fields)

    def test_only_actual_preferred_dividends_reduce_eps(self):
        f = facts("HK")
        c = cost(f)
        revenue = {y:100 for y in f["hist_periods"]+f["forecast_periods"]}
        before = foreign_pnl(f,c,revenue)["2026E"]["eps"]
        c["preferred_dividends"] = {y:1 for y in f["forecast_periods"]}
        c = require_cost(c,f)
        self.assertAlmostEqual(before-foreign_pnl(f,c,revenue)["2026E"]["eps"],.1)

    def test_exact_real_fields_are_parsed_without_other_comprehensive_income(self):
        rows = replies("HK")
        for row in rows[0]["data"]:
            row["indicators"] += [{"name":"财务费用","value":"0亿元人民币","source":"利润表"},
                {"name":"其他综合收益","value":"999亿元人民币","source":"利润表"},
                {"name":"应占联营公司损益","value":"1亿元人民币","source":"利润表"}]
        f = decode_financials(rows,ticker="hk01234",company="样例")
        self.assertEqual(f["income"]["财务费用"],[0,0,0])
        self.assertEqual(f["income"]["联营公司损益"],[1,1,1])
        self.assertEqual(f["income"]["其他经营净收益"],[None]*3)


if __name__ == "__main__":
    unittest.main()
