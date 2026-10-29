from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from valuation.historical_financials.foreign import currency_of_unit, decode_financials, quantity, report_metadata
from valuation.operating_cost.schema import require_cost
from valuation.shared.markets import (expenses, foreign_pnl, label, market_of,
    register_labels, valuation_factor, validate_pe_basis)
from valuation.workbook.anchors import cell_ref, find_row_by_label, find_label_position
from valuation.workbook.build import _hist as write_hist, _cost as write_cost, _pnl as write_pnl
from valuation.workbook.build import _valuation, build_workbook
from valuation.workbook.certify import certify_workbook
from valuation.research_dossier.snapshot import snapshot_from_spec


def replies(market="US"):
    rows = []
    for y in (2023, 2024, 2025):
        values = {"营业总收入": "100亿美元", "营业成本": "40亿美元", "毛利润": "60亿美元",
                  "销售一般管理费用": "10亿美元", "研发费用": "5亿美元", "营业利润": "45亿美元",
                  "税前利润": "46亿美元", "所得税": "10亿美元", "净利润": "36亿美元",
                  "归母净利润": "37亿美元", "Basic EPS": "0", "币种": "USD",
                  "流动资产合计": "30亿美元", "流动负债合计": "20亿美元"}
        if market == "HK":
            values = {k: v.replace("亿美元", "亿元人民币") for k, v in values.items()}
            values.update(销售成本="40亿元人民币", 销售费用="6亿元人民币", 管理费用="4亿元人民币",
                          基本每股收益="0", 币种="CNY")
        rows.append({"stockCode": "usTEST" if market == "US" else "hk01234", "stockName": "样例",
                     "reportPeriod": f"{y}Q4", "reportDate": f"{y}-09-27", "type": "当年累计",
                     "indicators": [{"name": k, "value": v, "source": "利润表"} for k, v in values.items()]})
    return [{"code": "0", "data": rows}]


def facts(market="US"):
    f = decode_financials(replies(market), ticker="usTEST" if market == "US" else "hk01234", company="样例",
                         metadata={"profit_basis": {"净利润": "adjusted", "归母净利润": "adjusted"}} if market == "US" else {})
    f.update(trading_currency="USD" if market == "US" else "HKD")
    f["market"] = {"当前股价": 10, "当前总股本_百万股": 1000, "每交易单位普通股数": 1,
                   "报表币种兑交易币种": 1.1}
    return f


def cost(f):
    years = f["forecast_periods"]
    raw = {key: {year: value for year in years} for key, value in {
        "gross_margin": .6, "sga_ratio": .1, "sales_ratio": .06, "admin_ratio": .04,
        "rd_ratio": .05, "tax_rate": .2, "minority": -.1, "other_op": 0,
        "nonop_inc": 1}.items()}
    raw["rationale"] = "预测税率20%，股本沿用1000百万股快照；营业外收入1亿仅用于公式校验。"
    return require_cost(raw, f)


def integration_spec(market="US"):
    """纯公式校验输入，不代表任何真实公司的分部或预测。"""
    f = facts(market)
    f["company"] = ("美股" if market == "US" else "港股") + "公式校验样例"
    segments, forecasts = [], {}
    for name in ["校验分部一", "校验分部二"]:
        seg = {"name": name, "parent": "", "method": "收入增速法", "unit": f["unit"],
               "hist_revenue": [50., 50., 50.], "note": "[推算]仅用于校验公式。"}
        segments.append(seg)
        forecasts[name] = {"method": "收入增速法", "historical_data": {"分部收入": {y: 50. for y in f["hist_periods"]}},
            "final_forecast": {"收入增速": {y: .1 for y in f["forecast_periods"]}},
            "revenue_unit": f["unit"], "final_rationale": "公式校验假设：各年增长10%，不代表投资预测。"}
    return {"facts": f, "cost": cost(f), "split": {"final_segments": segments, "official_parents": [],
            "split_logic": "仅用于跨表公式联动验证。", "structure_change": "", "other_note": "", "source_files": ""},
            "forecasts": forecasts, "consensus": {"institutions": []},
            "consensus_canonical": {"institution_estimates": [], "revenue_unit": f["unit"], "eps_unit": f["eps_unit"], "eps_scope": f["eps_scope"]},
            "comps": {"core": [{"name": f"校验同业{i}", "ticker": f"TEST{i}", "pe_y1": float(10+i),
                "pe_ttm": 10., "mcap": 100., "currency": "USD", "note": "仅用于检验公式",
                "eps_scope": f["eps_scope"], "forecast_period": "2026E"} for i in range(4)],
                "pe_adjust": 1., "rationale": "公式校验用倍数，不代表真实同业。"},
            "summary_notes": {field: [{"title": "公式校验", "text": "用于检查模型输出的跨表联动，不代表投资研究结论。"}]
                              for field in ["thesis", "outlook", "risks"]}}


class ForeignFactsTests(unittest.TestCase):
    def test_chinese_dollar_presentation_currency(self):
        meta = report_metadata([{"title": "公司：2025年年報", "viewpoint":
            "該等合併財務報表以美元列報，美元為本公司的功能貨幣及本集團的列報貨幣。",
            "url": "https://example.com/annual-report"}], ticker="hk01234", fiscal_year="2025")
        self.assertEqual(meta["reporting_currency"], "USD")
        self.assertEqual(meta["sources"][0]["url"], "https://example.com/annual-report")

    def test_presentation_currency_is_not_functional_currency(self):
        meta = report_metadata([{"title": "公司：二零二五年年報", "text":
            'The functional currency is United States Dollars. The consolidated financial statements are presented in Renminbi ("RMB").'}], ticker="hk01234", fiscal_year="2025")
        self.assertEqual(meta["reporting_currency"], "CNY")

    def test_share_metadata_requires_matching_security_row(self):
        meta = report_metadata([{"title": "Form 8-K", "text":
            '<table><tr><td>Common Stock</td><td>TEST</td><td>NASDAQ</td></tr></table>'}], ticker="usTEST", fiscal_year="2025")
        self.assertEqual(meta["ordinary_per_quote"], 1)
        other = report_metadata([{"text": '<tr><td>Common Stock</td><td>OTHER</td></tr>'}], ticker="usTEST", fiscal_year="2025")
        self.assertNotIn("ordinary_per_quote", other)

    def test_units_parentheses_zero_and_currency(self):
        self.assertEqual(quantity("(123.45)亿美元"), -123.45)
        self.assertEqual(quantity("5百万美元"), .05)
        self.assertEqual(quantity("0", kind="eps"), 0)
        self.assertIsNone(quantity("42"))
        self.assertIsNone(quantity("NaN"))

    def test_minority_negative_missing_asset_and_zero_eps(self):
        f = facts()
        self.assertEqual(f["income"]["少数股东损益"], [-1, -1, -1])
        self.assertEqual(f["income"]["EPS"], [0, 0, 0])
        self.assertEqual(f["balance"]["存货"], [None]*3)
        self.assertEqual(f["balance"]["营运资本"], [10]*3)
        self.assertEqual(f["income"]["历史加权平均股本_百万股"], [None]*3)

    def test_direct_minority_wins(self):
        r = replies()
        r[0]["data"][0]["indicators"].append({"name": "少数股东损益", "value": "2亿美元", "source": "利润表"})
        f = decode_financials(r, ticker="usTEST", company="样例")
        self.assertEqual(f["income"]["少数股东损益"][0], 2)

    def test_unknown_adjustment_scope_is_not_called_minority(self):
        f = decode_financials(replies(), ticker="usTEST", company="样例")
        self.assertEqual(f["income"]["少数股东损益"], [None]*3)
        self.assertTrue(any("调整范围" in x for x in f["notes"]["profit_checks"]))

    def test_adjusted_parent_alias_only_confirms_its_own_period(self):
        r=replies()
        r[0]["data"][-1]["indicators"].append({"name":"扣非归母净利润","value":"35亿美元","source":"利润表"})
        f=decode_financials(r,ticker="usTEST",company="样例")
        self.assertEqual(f["income"]["少数股东损益"],[None,None,1.])

    def test_basic_only_not_generic_eps_or_diluted(self):
        r = replies()
        for row in r[0]["data"]:
            row["indicators"] += [{"name": "EPS", "value": "999", "source": "利润表"},
                                  {"name": "稀释每股收益", "value": "888", "source": "利润表"}]
        self.assertEqual(decode_financials(r, ticker="usTEST", company="样例")["income"]["EPS"], [0]*3)

    def test_partial_optional_year_is_preserved(self):
        r = replies()
        r[0]["data"][0]["indicators"].append({"name": "存货合计", "value": "2亿美元", "source": "资产负债表"})
        self.assertEqual(decode_financials(r, ticker="usTEST", company="样例")["balance"]["存货"], [2, None, None])

    def test_hk_sales_cost_and_rd_in_admin(self):
        f = facts("HK")
        self.assertEqual(f["income"]["营业成本"], [40]*3)
        f["rd_in_admin"] = True
        p = foreign_pnl(f, cost(f), {y: 100 for y in f["hist_periods"]+f["forecast_periods"]})
        self.assertEqual(p["2026E"]["ebit"], 50)
        self.assertEqual(label(f, "研发费用"), "其中：研发费用")

    def test_fiscal_year_is_not_calendar_year_and_current_year_included(self):
        r = replies()
        r[0]["data"][-1].update(reportPeriod="2026Q4", reportDate="2026-01-25")
        r[0]["data"][0].update(reportPeriod="2025Q4", reportDate="2025-01-25")
        f = decode_financials(r, ticker="usTEST", company="样例")
        self.assertEqual(f["hist_periods"], ["2024A", "2025A", "2026A"])
        self.assertEqual(f["report_dates"]["2026A"], "2026-01-25")

    def test_mixed_currencies_rejected_and_other_company_filtered(self):
        r = replies()
        other = deepcopy(r[0]["data"][0]);other["stockCode"] = "usOTHER"
        other["indicators"][0]["value"] = "9亿港元"
        r[0]["data"].append(other)
        self.assertEqual(decode_financials(r,ticker="usTEST",company="样例")["currency"],"USD")
        r[0]["data"][0]["indicators"][0]["value"] = "9亿港元"
        with self.assertRaises(ValueError):decode_financials(r,ticker="usTEST",company="样例")

    def test_operating_residual_is_other_operating_income(self):
        for market, ticker, unit in [("HK", "hk01234", "亿元人民币"), ("US", "usTEST", "亿美元")]:
            with self.subTest(market=market):
                rows = replies(market)
                profit = "营业利润"
                for row in rows[0]["data"]:
                    for item in row["indicators"]:
                        if item["name"] == profit:
                            item["value"] = item["value"].replace("45", "44")
                facts = decode_financials(rows, ticker=ticker, company="样例")
                self.assertEqual(facts["income"]["其他经营净收益"], [-1, -1, -1])
                self.assertTrue(any("其他经营净收益=营业利润" in item for item in facts["derived"]))
                rows = replies(market)
                for row in rows[0]["data"]:
                    row["indicators"].append({"name": "其他经营净收益", "value": f"2{unit}", "source": "利润表"})
                kept = decode_financials(rows, ticker=ticker, company="样例")
                self.assertEqual(kept["income"]["其他经营净收益"], [2, 2, 2])
                self.assertFalse(any("其他经营净收益=" in item for item in kept["derived"]))

    def test_unknown_currency_not_inferred_from_listing(self):
        r = replies("HK")
        for row in r[0]["data"]:
            row["indicators"] = [dict(i, value=i["value"].replace("亿元人民币", "亿")) for i in row["indicators"] if i["name"]!="币种"]
        self.assertIsNone(decode_financials(r,ticker="hk01234",company="样例")["currency"])

    def test_row_unit_sets_currency_when_amounts_omit_it(self):
        r = replies("HK")
        for row in r[0]["data"]:
            row["unit"] = "美元-USD"
            row["indicators"] = [dict(i, value=i["value"].replace("亿元人民币", "亿")) for i in row["indicators"] if i["name"] != "币种"]
        f = decode_financials(r, ticker="hk01234", company="样例")
        self.assertEqual(f["currency"], "USD")
        self.assertEqual(f["unit"], "亿美元")
        self.assertEqual(f["income"]["营业收入"], [100, 100, 100])
        self.assertIsNone(currency_of_unit("欧元-EUR"))
        self.assertIsNone(currency_of_unit(""))

    def test_row_unit_conflict_is_rejected(self):
        r = replies()
        r[0]["data"][0]["unit"] = "港元-HKD"
        with self.assertRaises(ValueError):
            decode_financials(r, ticker="usTEST", company="样例")


class ForeignForecastTests(unittest.TestCase):
    def test_chain_and_expense_sensitivity(self):
        f=facts();c=cost(f);rev={y:100 for y in f["hist_periods"]+f["forecast_periods"]}
        before=foreign_pnl(f,c,rev)["2026E"]
        self.assertAlmostEqual(before["eps"], 4.048)
        c["rd_ratio"]["2026E"] += .01
        after=foreign_pnl(f,c,rev)["2026E"]
        self.assertAlmostEqual(before["eps"]-after["eps"], .088)
        f["market"]["当前总股本_百万股"] *= 2
        self.assertAlmostEqual(foreign_pnl(f,c,rev)["2026E"]["eps"],after["eps"]/2)

    def test_no_fake_a_share_fields_required_and_invalid_shares_rejected(self):
        f=facts();c=cost(f)
        self.assertNotIn("fin_exp",c)
        self.assertNotIn("sales_ratio",c)
        f["market"]["当前总股本_百万股"]=0
        with self.assertRaises(ValueError):foreign_pnl(f,c,{y:100 for y in f["hist_periods"]+f["forecast_periods"]})
        with self.assertRaises(Exception):require_cost(c,f)

    def test_missing_expense_does_not_become_zero(self):
        f=facts("HK");f["income"]["研发费用"]=[None]*3
        self.assertNotIn(("研发费用","rd_ratio"),expenses(f))
        c=cost(f)
        self.assertNotIn("rd_ratio",c)

    def test_fx_and_ads_applied_once(self):
        f=facts();f["currency"]="CNY";f["market"]["每交易单位普通股数"]=8
        self.assertAlmostEqual(valuation_factor(f),8.8)
        f["market"].pop("报表币种兑交易币种")
        with self.assertRaises(ValueError):valuation_factor(f)

    def test_supplier_forward_pe_survives_missing_basis(self):
        f=facts()
        validate_pe_basis(f,{"core":[{"name":"同业","pe_y1":10}]})
        validate_pe_basis(f,{"core":[{"name":"同业","pe_y1":10,"eps_scope":"SBCOC","forecast_period":"2026E"}]})

    def test_consensus_requires_scope_currency_and_ordinary_share_unit(self):
        from valuation.peer_valuation.fetch import parse_institution_estimates
        f=facts()
        block={"standardizedInfo":{"currency":"USD", "epsScope":"SBCOC", "epsUnit":"USD/普通股"},
               "forecastList":[{"orgNameAbbr":"测试机构", "publishDate":"2026-10-20",
                                "forecastItems":[{"year":2026,"revenue":120,"unit":"亿美元","basicEps":0}]}]}
        raw={"data":{"TEST":block}}
        houses,_=parse_institution_estimates(raw,f,shares=1000)
        self.assertEqual(houses[0]["estimates"]["2026E"]["eps"],0)
        block["standardizedInfo"]["epsUnit"]="USD/ADS"
        houses,_=parse_institution_estimates(raw,f,shares=1000)
        self.assertIsNone(houses[0]["estimates"]["2026E"]["eps"])
        block["standardizedInfo"]["currency"]="CNY"
        houses,_=parse_institution_estimates(raw,f,shares=1000)
        self.assertIsNone(houses[0]["estimates"]["2026E"]["revenue"])

    def test_explicit_ads_eps_is_not_labeled_as_ordinary(self):
        self.assertIsNone(quantity("8USD/ADS",kind="eps"))
        self.assertEqual(quantity("100百万股",kind="shares"),100)

    def test_missing_subject_scope_cannot_match_missing_peer_scope(self):
        f=facts();f["eps_scope"]=None
        with self.assertRaises(ValueError):
            validate_pe_basis(f,{"core":[{"name":"未知口径","pe_y1":10,"forecast_period":"2026E"}]})

    def test_market_detection_preserves_a_share(self):
        for s in ["300308","sh600519","sz000001"]:self.assertEqual(market_of(s),"A")
        for s in ["09992.HK","hk00700","09992"]:self.assertEqual(market_of(s),"HK")
        for s in ["AAPL","usBABA","BRK.B","HKD"]:self.assertEqual(market_of(s),"US")

    def test_rd_in_sga_is_not_deducted_twice(self):
        f=facts();f["rd_in_sga"]=True
        result=foreign_pnl(f,cost(f),{y:100 for y in f["hist_periods"]+f["forecast_periods"]})
        self.assertEqual(result["2026E"]["ebit"],50)


class ForeignWorkbookTests(unittest.TestCase):
    def test_all_seven_sheets_and_snapshot_pass_certification(self):
        for market in ["HK", "US"]:
            with self.subTest(market=market), tempfile.TemporaryDirectory() as directory:
                spec=integration_spec(market)
                with patch("valuation.workbook.build.load_spec",return_value=spec):
                    path=build_workbook(Path(directory))
                with patch("valuation.workbook.certify.load_spec",return_value=spec):
                    report=certify_workbook(Path(directory),path)
                self.assertEqual(report["calc_fails"],[])
                self.assertEqual(report["format_fails"],[])
                snapshot=snapshot_from_spec(spec)
                self.assertGreater(snapshot["valuation"]["target_price"],0)

    def test_labels_formulas_missing_values_and_conversion(self):
        f=facts();c=cost(f);wb=Workbook();wb.remove(wb.active)
        for n in ["历史财务数据","收入预测","运营成本预测","利润表预测","估值预测"]:wb.create_sheet(n)
        register_labels(wb,f)
        write_hist(wb,{"facts":f})
        rev=wb["收入预测"];periods=f["hist_periods"]+f["forecast_periods"]
        for i,y in enumerate(periods,2):rev.cell(2,i,y);rev.cell(3,i,100)
        rev.cell(3,1,"营业收入")
        write_cost(wb,{"facts":f,"cost":c})
        write_pnl(wb,{"facts":f,"cost":c})
        fs,pnl=wb["历史财务数据"],wb["利润表预测"]
        self.assertEqual(fs.cell(find_row_by_label(fs,"少数股东损益"),1).value,"少数股东损益")
        r,c=find_label_position(fs,"存货")
        self.assertEqual(c,6)
        self.assertEqual(fs.cell(r,c+1).value,"-")
        self.assertEqual(fs.cell(find_row_by_label(fs,"EPS"),1).value,"基本每股收益（扣非）")
        self.assertEqual(pnl.cell(find_row_by_label(pnl,"营业利润"),1).value,"营业利润（SOPR）")
        eps=pnl.cell(find_row_by_label(pnl,"EPS"),5).value
        self.assertIn(cell_ref(pnl,"归母净利润","2026E",same_sheet=True),eps)
        self.assertIn(f"'历史财务数据'!B{find_row_by_label(fs,'当前总股本_百万股')}",eps)
        self.assertFalse(any("预测基本加权平均普通股数" in str(cell.value) for row in pnl for cell in row))
        self.assertEqual(fs.max_column,9)
        self.assertEqual(fs["F2"].value,"资产负债表（亿美元）")


if __name__=="__main__":unittest.main()
