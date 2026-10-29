"""以 A 股模板为公共格式契约，防止港美再次另建布局。"""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from valuation.tests.test_foreign_financials import integration_spec
from valuation.workbook.build import build_workbook, SHEET_ORDER
from valuation.workbook.anchors import cell_ref, find_col_by_header, find_row_by_label
from valuation.shared.markets import register_labels


class CommonMarketTemplateTests(unittest.TestCase):
    def build(self, market, directory, *, rd_subset=False):
        spec=deepcopy(integration_spec(market))
        if rd_subset:spec['facts']['rd_in_sga']=True
        with patch('valuation.workbook.build.load_spec',return_value=spec):
            path=build_workbook(Path(directory))
        wb=load_workbook(path);register_labels(wb,spec['facts'])
        return wb,spec

    def test_hk_and_us_share_a_share_sections_and_period_columns(self):
        for market in ['HK','US']:
            with self.subTest(market=market),tempfile.TemporaryDirectory() as directory:
                wb,spec=self.build(market,directory)
                self.assertEqual(wb.sheetnames,SHEET_ORDER)
                self.assertEqual(wb['历史财务数据'].max_column,9)
                self.assertIn('利润表',wb['历史财务数据']['A2'].value)
                self.assertIn('资产负债表',wb['历史财务数据']['F2'].value)
                self.assertEqual(wb['历史财务数据'].column_dimensions['E'].width,3)
                for name in ['运营成本预测','利润表预测']:
                    ws=wb[name]
                    self.assertEqual(ws.max_column,8)
                    self.assertEqual(ws['H2'].value,'备注')
                    self.assertEqual([ws.cell(2,c).value for c in range(2,8)],
                                     spec['facts']['hist_periods']+spec['facts']['forecast_periods'])
                    self.assertEqual(ws.column_dimensions['A'].width,32)
                cost=wb['运营成本预测']
                for title in ['成本毛利区','费用区','总费用']:
                    self.assertGreater(find_row_by_label(cost,title),2)
                pnl=wb['利润表预测']
                for title in ['营业收入同比','营业成本','合并毛利率','归母净利润同比']:
                    self.assertGreater(find_row_by_label(pnl,title),2)
                summary=wb['总结']
                self.assertEqual(summary.max_column,11)
                self.assertEqual(summary.column_dimensions['F'].width,3)
                for left,right in [('估值概览','当前行情'),('财务预测','关键经营假设'),('投资逻辑','未来展望')]:
                    self.assertEqual(find_row_by_label(summary,left),find_row_by_label(summary,right))
                self.assertTrue(any('手动输入' in str(c.value) for row in summary for c in row))
                self.assertFalse(any('· EPS' in str(ws['A1'].value) for ws in wb))

    def test_long_rationale_stays_inside_notes_column(self):
        with tempfile.TemporaryDirectory() as directory:
            wb,spec=self.build('HK',directory)
            ws=wb['运营成本预测']
            self.assertEqual(ws['H3'].value,spec['cost']['rationale'])
            self.assertTrue(ws['H3'].alignment.wrap_text)
            self.assertEqual(ws['H3'].alignment.vertical,'top')
            self.assertTrue(any(m.min_col==8 and m.max_col==8 and m.min_row==3 for m in ws.merged_cells.ranges))
            self.assertEqual(ws.max_column,8)

    def test_empty_consensus_and_zero_base_growth_are_explicitly_missing(self):
        from valuation.workbook.build import consensus_median_formula
        from valuation.workbook.market_config import ratio_formula
        self.assertEqual(consensus_median_formula(3,20,'C',guarded=True),
                         '=IF(COUNT(C3:C20)>0,MEDIAN(C3:C20),"-")')
        self.assertEqual(consensus_median_formula(3,20,'C'),'=MEDIAN(C3:C20)')
        self.assertIn('B3<>0',ratio_formula('C3','B3',growth=True,guarded=True))
        self.assertEqual(ratio_formula('C3','B3',growth=True),'=C3/B3-1')

    def test_expenses_flow_from_cost_build_and_rd_subset_is_not_double_counted(self):
        with tempfile.TemporaryDirectory() as directory:
            wb,spec=self.build('US',directory,rd_subset=True)
            year=spec['facts']['forecast_periods'][0]
            pnl,cost=wb['利润表预测'],wb['运营成本预测']
            pcol=find_col_by_header(pnl,year);ccol=find_col_by_header(cost,year)
            sga=pnl.cell(find_row_by_label(pnl,'SG&A'),pcol).value
            self.assertIn(cell_ref(cost,'SG&A',year),sga)
            total=cost.cell(find_row_by_label(cost,'总费用'),ccol).value
            self.assertIn(cell_ref(cost,'SG&A',year,same_sheet=True),total)
            self.assertNotIn(cell_ref(cost,'研发费用',year,same_sheet=True),total)
            op=pnl.cell(find_row_by_label(pnl,'营业利润'),pcol).value
            self.assertNotIn(cell_ref(pnl,'研发费用',year,same_sheet=True),op)


if __name__=='__main__':unittest.main()
