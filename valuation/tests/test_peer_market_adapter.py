"""同业查询带市场，最新预测年PE与备注沿用A股逻辑。"""
from copy import deepcopy
from datetime import date
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from valuation.peer_valuation.fetch import (_candidate_query, _fill_one, _norm_ticker,
    fill_candidates, fetch_subject_consensus, parse_institution_estimates)
from valuation.peer_valuation.schema import require_comps
from valuation.shared.markets import validate_pe_basis
from valuation.tests.test_foreign_financials import facts, integration_spec
from valuation.workbook.build import build_workbook
from valuation.workbook.certify import certify_workbook
from valuation.workbook.adapt import adapt_consensus
from valuation.research_dossier.snapshot import snapshot_from_spec
from valuation.tests.test_valuation_fx import reply as fx_reply


def quote(query, code, pe="16.5", **metadata):
    return {"query":query,"standardized_info":{"stockFullCode":code,"stockCode":code.removeprefix('hk'),"currency":"HKD"},
            "valuation":{"peForward":pe,**metadata}}


class PeerMarketAdapterTests(unittest.TestCase):
    def test_queries_respect_candidate_market(self):
        f=facts("HK")
        cases=[({"ticker":"2122"},"hk02122"),({"ticker":"0150"},"hk00150"),
               ({"ticker":"0700","market":"HK"},"hk00700"),({"ticker":"700.HK"},"hk00700"),
               ({"ticker":"hk700"},"hk00700"),({"ticker":"8136","market":"JP"},"jpn8136"),
               ({"ticker":"8136.T"},"jpn8136"),({"ticker":"jpn8136"},"jpn8136"),
               ({"ticker":"002292","market":"A"},"002292"),({"ticker":"MAT","market":"US"},"usMAT"),
               ({"ticker":"BRK.B","market":"US"},"usBRK.B")]
        for item,expected in cases:
            with self.subTest(item=item):self.assertEqual(_candidate_query(item,f),expected)
        self.assertEqual(_candidate_query({"ticker":"8136"},facts("US")),"")
        self.assertEqual(_norm_ticker("hk700"),_norm_ticker("00700.HK"))
        self.assertNotEqual(_norm_ticker("hk02122"),_norm_ticker("jpn2122"))

    def test_missing_basis_does_not_erase_returned_forward_pe(self):
        mcp=Mock()
        mcp.financial_snapshot.return_value={"data":{}}
        mcp.price_performance.return_value={"data":[quote("hk09896","hk09896")]}
        with tempfile.TemporaryDirectory() as directory:
            filled=fill_candidates(mcp,[{"name":"名创优品","ticker":"09896"}],facts("HK"),Path(directory),[])
        self.assertEqual(filled[0]["pe_y1"],16.5)
        self.assertFalse(filled[0]["thin"])
        self.assertNotIn("basis_gap",filled[0])
        self.assertIsNone(filled[0]["eps_scope"])
        mcp.price_performance.assert_called_once()
        self.assertEqual(mcp.price_performance.call_args.args[0],["hk09896"])

    def test_wrong_market_response_cannot_supply_pe_or_market_cap(self):
        row={"ticker":"02122","query":"hk02122","name":"Kidsland"}
        wrong_snapshot={"standardizedInfo":{"fullCode":"jpn2122","ticker":"2122"},
                        "snapshotItems":[{"name_cn":"总市值","value":999999999}]}
        wrong_price=quote("hk02122","jpn2122","25")
        item=_fill_one(row,wrong_snapshot,wrong_price)
        self.assertIsNone(item["pe_y1"])
        self.assertIsNone(item["mcap"])

    def test_known_conflicting_basis_and_year_remain_rejected(self):
        f=facts("HK")
        for fields in [{"eps_scope":"SBCOC"},{"forecast_period":"2027E"}]:
            with self.subTest(fields=fields),self.assertRaises(ValueError):
                validate_pe_basis(f,{"core":[{"name":"同业","pe_y1":10,**fields}]})

    def test_core_card_and_seven_sheet_valuation_accept_supplier_pe(self):
        spec=deepcopy(integration_spec("HK"))
        filled=spec["comps"]["core"]
        for row in filled:
            row.pop("eps_scope",None);row.pop("forecast_period",None)
            row["pe_source"]="pricePerformance.valuation.peForward"
        spec["comps"]=require_comps(spec["comps"],spec["facts"],filled)
        self.assertEqual(spec["comps"]["open_gaps"],[])
        snapshot=snapshot_from_spec(spec)
        self.assertGreater(snapshot["valuation"]["target_price"],0)
        self.assertEqual(snapshot["valuation"]["pe_basis"],"2026E PE")
        self.assertNotIn("basis_warnings",snapshot["valuation"])
        with tempfile.TemporaryDirectory() as directory:
            with patch("valuation.workbook.build.load_spec",return_value=spec):path=build_workbook(Path(directory))
            with patch("valuation.workbook.certify.load_spec",return_value=spec):report=certify_workbook(Path(directory),path)
            self.assertEqual(report["calc_fails"],[])
            self.assertEqual(report["format_fails"],[])
            wb=load_workbook(path)
            ws=wb["估值预测"]
            self.assertTrue(any(c.value=="2026E PE" for row in ws for c in row))
            self.assertFalse(any("未返回盈利口径" in str(c.value) for row in ws for c in row))

    def test_provider_name_cannot_be_overridden_by_a_wrong_candidate_name(self):
        row={"ticker":"003025","query":"003025","name":"布鲁可"}
        price={"standardized_info":{"stockFullCode":"sz003025","stockShortName":"思进智能","currency":"CNY"},
               "valuation":{"peForward":"22.32"}}
        item=_fill_one(row,{},price)
        self.assertEqual(item["name"],"思进智能（003025）")
        self.assertNotIn("布鲁可",item["name"])
        self.assertIsNone(item["pe_y1"])
        self.assertTrue(item["thin"])
        self.assertIn("不匹配",item["data_gap"])


class ConsensusRevenueTests(unittest.TestCase):
    def raw(self, *lines):
        return {"data": {"subject": {"forecastList": [{"orgNameAbbr": "机构一",
            "publishDate": "2026-10-22", "forecastItems": list(lines)}]}}}

    def test_unit_currency_restores_revenue_without_requiring_eps(self):
        for market, currency in [("HK", "CNY"), ("US", "USD")]:
            with self.subTest(market=market):
                f = facts(market)
                raw = self.raw({"year": 2026, "revenue": 35_557_000_000, "unit": currency},
                               {"year": 2027, "revenue": 0, "unit": currency})
                houses, _ = parse_institution_estimates(raw, f, shares=1_000_000_000)
                values = houses[0]["estimates"]
                self.assertEqual(values["2026E"]["revenue"], 355.57)
                self.assertEqual(values["2027E"]["revenue"], 0)
                self.assertIsNone(values["2026E"]["eps"])
                self.assertIsNone(values["2028E"]["revenue"])
                # 机构逐年预测币种优先于公司层面的交易币种。
                raw["data"]["subject"]["currency"] = "HKD"
                houses, _ = parse_institution_estimates(raw, f, shares=None)
                self.assertEqual(houses[0]["estimates"]["2026E"]["revenue"], 355.57)

    def test_conversion_uses_verified_fx_in_both_directions(self):
        for report_currency, source_currency, expected in [("CNY", "HKD", 85.63), ("HKD", "CNY", 100/0.8563)]:
            with self.subTest(report_currency=report_currency):
                f = facts("HK")
                f["currency"] = report_currency
                f["valuation_fx"] = {"base": "CNY", "quote": "HKD", "rate": 1/0.8563,
                    "date": "2026-10-27", "source": "中国外汇交易中心"}
                raw = self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": source_currency})
                houses, _ = parse_institution_estimates(raw, f, shares=None)
                value = houses[0]["estimates"]["2026E"]
                self.assertAlmostEqual(value["revenue"], expected)
                self.assertEqual(value["revenue_source"]["currency"], source_currency)
                self.assertEqual(value["revenue_source"]["fx"]["quote"], report_currency)

    def test_unknown_currency_or_missing_fx_stays_missing(self):
        for currency in ["HKD", "USD", "亿元", None]:
            with self.subTest(currency=currency):
                houses, gaps = parse_institution_estimates(
                    self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": currency}), facts("HK"), shares=None)
                self.assertIsNone(houses[0]["estimates"]["2026E"]["revenue"])
                self.assertTrue(any("收入币种" in gap for gap in gaps))
        for rate in [0, -1, float("nan"), float("inf")]:
            with self.subTest(rate=rate):
                f = facts("HK")
                f["valuation_fx"] = {"base": "CNY", "quote": "HKD", "rate": rate}
                houses, _ = parse_institution_estimates(
                    self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": "HKD"}), f, shares=None)
                self.assertIsNone(houses[0]["estimates"]["2026E"]["revenue"])

    def test_explicit_currency_precedes_unit_and_eps_rules_are_unchanged(self):
        f = facts("HK")
        line = {"year": 2026, "revenue": 355.57, "unit": "亿元", "currency": "CNY",
                "basicEps": 2, "epsScope": f["eps_scope"], "epsUnit": f["eps_unit"]}
        houses, _ = parse_institution_estimates(self.raw(line), f, shares=None)
        value = houses[0]["estimates"]["2026E"]
        self.assertEqual(value["revenue"], 355.57)
        self.assertEqual(value["eps"], 2)
        del line["epsScope"]
        houses, _ = parse_institution_estimates(self.raw(line), f, shares=None)
        self.assertIsNone(houses[0]["estimates"]["2026E"]["eps"])
        self.assertEqual(houses[0]["estimates"]["2026E"]["revenue"], 355.57)

    def test_revenue_median_and_workbook_adapter_receive_restored_values(self):
        f = facts("HK")
        raw = self.raw({"year": 2026, "revenue": 35_557_000_000, "unit": "CNY"})
        block = raw["data"]["subject"]
        for name, amount in [("机构二", 33_817_000_000), ("机构三", 33_393_000_000)]:
            block["forecastList"].append({"orgNameAbbr": name, "publishDate": "2026-10-22",
                "forecastItems": [{"year": 2026, "revenue": amount, "unit": "CNY"}]})
        block["consensusSummary"] = [{"year": 2026, "revenueMedian": 999_000_000_000, "unit": "HKD"}]
        mcp = Mock()
        mcp.profit_forecast.return_value = raw
        with tempfile.TemporaryDirectory() as directory:
            cons = fetch_subject_consensus(mcp, f, Path(directory), [])
        self.assertEqual(cons["consensus"]["2026E"]["revenue"]["value"], 338.17)
        self.assertEqual(cons["consensus"]["2026E"]["revenue"]["sample_count"], 3)
        self.assertIsNone(cons["consensus"]["2026E"]["eps"]["value"])
        adapted = adapt_consensus(cons, f)
        self.assertEqual(adapted["institutions"][0]["revenue"], [355.57, None, None])
        self.assertIn(f["unit"], adapted["revenue_note"])

    def test_a_share_revenue_and_eps_follow_original_path(self):
        f = {"ticker": "300308", "forecast_periods": ["2026E"]}
        raw = self.raw({"year": 2026, "revenue": 10_000_000_000, "parentNetProfit": 2_000_000_000, "unit": "CNY"})
        houses, _ = parse_institution_estimates(raw, f, shares=1_000_000_000)
        self.assertEqual(houses[0]["estimates"]["2026E"], {"revenue": 100., "eps": 2.})


class LiveConsensusFxTests(unittest.TestCase):
    raw = ConsensusRevenueTests.raw

    def setUp(self):
        for module in ("valuation.peer_valuation.fetch", "valuation.historical_financials.fx"):
            patcher = patch(module + ".run_date", return_value=date(2026, 10, 28))
            patcher.start()
            self.addCleanup(patcher.stop)

    def client(self, raw):
        mcp = Mock()
        mcp.profit_forecast.return_value = raw
        def get_rate(tool, arguments):
            self.assertEqual(tool, "get_economic_indicator_data")
            self.assertEqual(arguments["data_source"], "中国外汇交易中心")
            name = arguments["index_name"]
            currency = "HKD" if "港元" in name else "USD"
            value = .85 if currency == "HKD" else 7
            return fx_reply(currency, [("2026-10-28", value), ("2026-10-27", value)])
        mcp.call.side_effect = get_rate
        return mcp

    def test_all_six_currency_directions_across_a_hk_us(self):
        for market, target, source, rate in [
            ("A", "CNY", "USD", 7), ("A", "CNY", "HKD", .85),
            ("HK", "HKD", "CNY", 1/.85), ("HK", "HKD", "USD", 7/.85),
            ("US", "USD", "CNY", 1/7), ("US", "USD", "HKD", .85/7),
        ]:
            with self.subTest(market=market, target=target, source=source):
                f = facts(market) if market != "A" else {"ticker": "300308", "market_type": "A",
                    "forecast_periods": ["2026E"], "market": {"当前总股本_百万股": 1000}}
                f["currency"] = target
                f["unit"] = {"CNY": "亿元人民币", "HKD": "亿港元", "USD": "亿美元"}[target]
                # 已有目标价汇率与机构营收转换无关，不能覆盖本次查询。
                f["valuation_fx"] = {"base": source, "quote": target, "rate": 999}
                before = deepcopy(f)
                raw = self.raw({"year": 2026, "revenue": 10_000_000_000,
                                "parentNetProfit": 1_000_000_000, "unit": source})
                mcp = self.client(raw)
                with tempfile.TemporaryDirectory() as directory:
                    cons = fetch_subject_consensus(mcp, f, Path(directory), [])
                estimate = cons["institution_estimates"][0]["estimates"]["2026E"]
                self.assertAlmostEqual(estimate["revenue"], 100 * rate)
                self.assertEqual(estimate["revenue_source"]["fx"]["date"], "2026-10-28")
                self.assertEqual(f, before)
                self.assertEqual(mcp.call.call_count, 2 if "CNY" not in (source, target) else 1)
                self.assertIn("2026-10-28", adapt_consensus(cons, f)["revenue_note"])

    def test_equal_model_and_trading_currency_still_fetches_for_forecast(self):
        f = facts("US")
        self.assertEqual(f["currency"], f["trading_currency"])
        mcp = self.client(self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": "CNY"}))
        with tempfile.TemporaryDirectory() as directory:
            cons = fetch_subject_consensus(mcp, f, Path(directory), [])
        self.assertAlmostEqual(cons["consensus"]["2026E"]["revenue"]["value"], 100/7)
        mcp.call.assert_called_once()

    def test_missing_model_currency_is_an_upstream_error(self):
        f = facts("HK")
        f.pop("currency")
        mcp = self.client(self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": "USD"}))
        with tempfile.TemporaryDirectory() as directory, self.assertRaisesRegex(ValueError, "模型财务币种尚未确定"):
            fetch_subject_consensus(mcp, f, Path(directory), [])
        mcp.call.assert_not_called()

    def test_same_currency_needs_no_fx_and_ignores_aggregate_currency(self):
        f = facts("HK")
        raw = self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": "CNY"})
        raw["data"]["subject"]["currency"] = "HKD"
        raw["data"]["subject"]["consensusSummary"] = [{"year": 2026, "unit": "HKD", "revenueMedian": 999}]
        mcp = self.client(raw)
        with tempfile.TemporaryDirectory() as directory:
            cons = fetch_subject_consensus(mcp, f, Path(directory), [])
        self.assertEqual(cons["consensus"]["2026E"]["revenue"]["value"], 100)
        mcp.call.assert_not_called()

    def test_multiple_forecast_currencies_use_same_previous_day(self):
        f = facts("HK")
        raw = self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": "USD"},
                       {"year": 2027, "revenue": 10_000_000_000, "unit": "HKD"})
        mcp = self.client(raw)
        mcp.call.side_effect = [fx_reply("HKD", [("2026-10-28", .84), ("2026-10-27", .85)]),
                               fx_reply("USD", [("2026-10-27", 7)])]
        with tempfile.TemporaryDirectory() as directory:
            cons = fetch_subject_consensus(mcp, f, Path(directory), [])
        estimates = cons["institution_estimates"][0]["estimates"]
        self.assertEqual(estimates["2026E"]["revenue"], 700)
        self.assertEqual(estimates["2027E"]["revenue"], 85)
        self.assertEqual({estimates[y]["revenue_source"]["fx"]["date"] for y in ("2026E", "2027E")},
                         {"2026-10-27"})

    def test_missing_fx_cannot_reuse_old_valuation_rate_or_succeed_with_all_blanks(self):
        f = facts("HK")
        f["valuation_fx"] = {"base": "USD", "quote": "CNY", "rate": 999}
        mcp = self.client(self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": "USD"}))
        mcp.call.return_value = fx_reply("USD", [("2026-10-26", 7)])
        mcp.call.side_effect = None
        with tempfile.TemporaryDirectory() as directory, self.assertRaisesRegex(ValueError, "接口有预测金额"):
            fetch_subject_consensus(mcp, f, Path(directory), [])

    def test_unknown_forecast_currency_is_reported(self):
        for unit in (None, "亿元", "EUR"):
            with self.subTest(unit=unit):
                mcp = self.client(self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": unit}))
                with tempfile.TemporaryDirectory() as directory, self.assertRaisesRegex(ValueError, "收入币种"):
                    fetch_subject_consensus(mcp, facts("HK"), Path(directory), [])
                mcp.call.assert_not_called()

    def test_unsupported_detail_currency_cannot_use_parent_currency(self):
        raw = self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": "EUR"})
        raw["data"]["subject"]["currency"] = "CNY"
        mcp = self.client(raw)
        with tempfile.TemporaryDirectory() as directory, self.assertRaisesRegex(ValueError, "收入币种"):
            fetch_subject_consensus(mcp, facts("HK"), Path(directory), [])
        mcp.call.assert_not_called()

    def test_partial_failure_preserves_same_currency_forecast_and_gap(self):
        raw = self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": "CNY"},
                       {"year": 2027, "revenue": 10_000_000_000, "unit": "USD"})
        mcp = self.client(raw)
        mcp.call.side_effect = RuntimeError("test unavailable")
        with tempfile.TemporaryDirectory() as directory:
            cons = fetch_subject_consensus(mcp, facts("HK"), Path(directory), [])
        self.assertEqual(cons["consensus"]["2026E"]["revenue"]["value"], 100)
        self.assertIsNone(cons["consensus"]["2027E"]["revenue"]["value"])
        self.assertTrue(any("汇率获取失败" in gap for gap in cons["open_gaps"]))
        self.assertIn("部分机构营收", adapt_consensus(cons, facts("HK"))["revenue_note"])

    def test_absent_forecasts_are_unavailable_but_real_zero_survives(self):
        mcp = self.client(self.raw({"year": 2026, "revenue": None, "unit": "USD"}))
        with tempfile.TemporaryDirectory() as directory:
            cons = fetch_subject_consensus(mcp, facts("HK"), Path(directory), [])
        self.assertEqual(cons["coverage_status"], "unavailable")
        mcp.call.assert_not_called()
        mcp = self.client(self.raw({"year": 2026, "revenue": 0, "unit": "USD"}))
        with tempfile.TemporaryDirectory() as directory:
            cons = fetch_subject_consensus(mcp, facts("HK"), Path(directory), [])
        self.assertEqual(cons["consensus"]["2026E"]["revenue"]["value"], 0)

    def test_each_fetch_queries_again_and_ignores_unused_forecast_years(self):
        mcp = self.client(self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": "USD"},
                                  {"year": 2025, "revenue": 10_000_000_000, "unit": "HKD"}))
        with tempfile.TemporaryDirectory() as directory:
            fetch_subject_consensus(mcp, facts("HK"), Path(directory), [])
            fetch_subject_consensus(mcp, facts("HK"), Path(directory), [])
        self.assertEqual(mcp.call.call_count, 2)
        self.assertTrue(all("美元兑人民币" in c.args[1]["index_name"] for c in mcp.call.call_args_list))

    def test_converted_revenue_reaches_workbook_and_median(self):
        spec = integration_spec("HK")
        raw = self.raw(*[{"year": year, "revenue": 10_000_000_000, "unit": "USD"}
                         for year in (2026, 2027, 2028)])
        mcp = self.client(raw)
        with tempfile.TemporaryDirectory() as directory:
            cons = fetch_subject_consensus(mcp, spec["facts"], Path(directory), [])
            spec["consensus_canonical"] = cons
            spec["consensus"] = adapt_consensus(cons, spec["facts"])
            snapshot = snapshot_from_spec(spec)
            self.assertEqual(snapshot["consensus"]["2026E"]["revenue"], 700)
            with patch("valuation.workbook.build.load_spec", return_value=spec):
                path = build_workbook(Path(directory))
            wb = load_workbook(path)
            self.assertEqual([wb["估值预测"].cell(3, c).value for c in (3, 4, 5)], [700, 700, 700])
            self.assertIn("2026-10-28", spec["consensus"]["revenue_note"])

    def test_latest_report_is_selected_per_metric_and_forecast_year(self):
        raw = self.raw({"year": 2026, "revenue": 10_000_000_000, "unit": "USD"},
                       {"year": 2027, "revenue": 0, "unit": "USD"})
        raw["data"]["subject"]["forecastList"].append({"orgNameAbbr": "机构一",
            "publishDate": "2026-10-20", "forecastItems": [
                {"year": 2026, "revenue": 9_000_000_000, "unit": "USD"},
                {"year": 2027, "revenue": 8_000_000_000, "unit": "USD"},
                {"year": 2028, "revenue": 7_000_000_000, "unit": "USD"}]})
        mcp = self.client(raw)
        with tempfile.TemporaryDirectory() as directory:
            cons = fetch_subject_consensus(mcp, facts("US"), Path(directory), [])
        self.assertEqual(len(cons["institution_estimates"]), 1)
        house = cons["institution_estimates"][0]
        self.assertEqual([house["estimates"][y]["revenue"] for y in ("2026E", "2027E", "2028E")], [100, 0, 70])
        self.assertEqual(house["estimates"]["2028E"]["revenue_source"]["report_date"], "2026-10-20")
        self.assertEqual(cons["consensus"]["2028E"]["revenue"]["sample_count"], 1)
        self.assertIn("2028E采用2026-10-20报告", adapt_consensus(cons, facts("US"))["revenue_note"])

    def test_empty_new_report_does_not_replace_used_report_date(self):
        raw = self.raw({"year": 2026, "unit": "USD"})
        raw["data"]["subject"]["forecastList"].append({"orgNameAbbr": "机构一",
            "publishDate": "2026-10-20", "forecastItems": [{"year": 2026, "revenue": 9_000_000_000, "unit": "USD"}]})
        mcp = self.client(raw)
        with tempfile.TemporaryDirectory() as directory:
            cons = fetch_subject_consensus(mcp, facts("US"), Path(directory), [])
        self.assertEqual(cons["institution_estimates"][0]["report_date"], "2026-10-20")
        self.assertEqual(cons["consensus"]["2026E"]["revenue"]["value"], 90)


if __name__=="__main__":unittest.main()
