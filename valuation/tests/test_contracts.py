from __future__ import annotations

import unittest

from valuation.research_dossier.snapshot import _segment_roll
from valuation.segment_research.methods import formula_errors, roll_segment
from valuation.segment_research.schema import as_growth_rate, require_forecast, validate_forecast
from valuation.segment_split.hist_agent import (
    HistCell,
    HistFill,
    HistFillDeps,
    HistSeg,
    _note_tag,
    hist_errors,
    plan_slot_errors,
)
from valuation.segment_split.split_agent import SegmentPick, SplitAgentDeps, SplitPlan, _fill_plan
from valuation.segment_split.schema import t_recon
from valuation.segment_split.split_plan import is_residual_name, normalize_split_plan, validate_split_plan
from valuation.workbook.adapt import adapt_split
from valuation.workbook.certify import formula_row_labels, hist_recon_errors
from valuation.workbook.anchors import find_row_in_group
from valuation.workbook.tree import build_split_groups, top_level_revenue_labels


class ResidualNameTests(unittest.TestCase):
    def test_two_qita_become_unique_names(self):
        raw = {
            "caliber": "drilled",
            "official_parents": ["消费电子", "汽车电子", "通讯及数据中心"],
            "segments": [
                {"name": "消费电子", "parent": "消费电子", "method": "收入增速法", "why_method": "年报行"},
                {"name": "汽车电子", "parent": "汽车电子", "method": "收入增速法", "why_method": "年报行"},
                {"name": "电脑互联", "parent": "通讯及数据中心", "method": "收入增速法", "why_method": "下探"},
                {"name": "通讯互联", "parent": "通讯及数据中心", "method": "收入增速法", "why_method": "下探"},
                {"name": "其他", "parent": "", "method": "收入增速法", "why_method": "公司级残差"},
                {"name": "其他", "parent": "通讯及数据中心", "method": "收入增速法", "why_method": "父口径其余"},
            ],
        }
        plan = normalize_split_plan(raw)
        names = [item["name"] for item in plan["segments"]]
        self.assertIn("其他", names)
        self.assertIn("通讯及数据中心其他", names)
        self.assertEqual(len(names), len(set(names)))
        by_name = {item["name"]: item for item in plan["segments"]}
        self.assertEqual(by_name["其他"]["parent"], "")
        self.assertEqual(by_name["通讯及数据中心其他"]["parent"], "通讯及数据中心")
        for item in plan["segments"]:
            self.assertNotIn("method", item)
            self.assertNotIn("why_method", item)

    def test_official_plan_gets_company_residual(self):
        deps = SplitAgentDeps(
            company="宁德时代",
            ticker="300750",
            forecast_years=["2026"],
            official_names=["动力电池系统", "储能电池系统"],
            official_rows=[],
            search=lambda *_a, **_k: [],
        )
        plan = SplitPlan(
            caliber="official",
            official_parents=["动力电池系统", "储能电池系统"],
            segments=[
                SegmentPick(name="动力电池系统", parent="动力电池系统"),
                SegmentPick(name="储能电池系统", parent="储能电池系统"),
            ],
        )
        filled = _fill_plan(deps, plan)
        names = [item.name for item in filled.segments]
        self.assertIn("其他", names)
        other = next(item for item in filled.segments if item.name == "其他")
        self.assertEqual(other.parent, "")
        self.assertFalse(hasattr(other, "method") and other.method)

    def test_normalize_injects_company_residual(self):
        plan = normalize_split_plan(
            {
                "caliber": "official",
                "official_parents": ["动力电池系统", "储能电池系统"],
                "segments": [
                    {"name": "动力电池系统", "parent": "动力电池系统", "method": "收入增速法", "why_method": "年报行"},
                    {"name": "储能电池系统", "parent": "储能电池系统", "method": "收入增速法", "why_method": "年报行"},
                ],
            }
        )
        names = [item["name"] for item in plan["segments"]]
        self.assertIn("其他", names)
        other = next(item for item in plan["segments"] if item["name"] == "其他")
        self.assertEqual(other["parent"], "")
        self.assertNotIn("method", other)

    def test_validate_requires_company_residual(self):
        errors = validate_split_plan(
            {
                "caliber": "official",
                "official_parents": ["动力电池系统", "储能电池系统"],
                "segments": [
                    {"name": "动力电池系统", "parent": "动力电池系统", "method": "收入增速法", "why_method": "年报行"},
                    {"name": "储能电池系统", "parent": "储能电池系统", "method": "收入增速法", "why_method": "年报行"},
                ],
            }
        )
        self.assertTrue(any("公司级残差" in item for item in errors))


class MethodPickTests(unittest.TestCase):
    def test_split_normalize_does_not_keep_method(self):
        plan = normalize_split_plan(
            {
                "caliber": "official",
                "official_parents": ["动力电池系统", "储能电池系统"],
                "segments": [
                    {"name": "动力电池系统", "parent": "动力电池系统"},
                    {"name": "储能电池系统", "parent": "储能电池系统"},
                    {"name": "其他", "parent": ""},
                ],
            }
        )
        self.assertEqual(validate_split_plan(plan), [])
        for item in plan["segments"]:
            self.assertNotIn("method", item)
            self.assertNotIn("why_method", item)

    def test_validator_does_not_judge_method_from_brief(self):
        facts = _facts()
        notes = _qty_notes()
        self.assertEqual(validate_forecast(_growth_payload(0.18, 10.0), facts, notes), [])
        card = require_forecast(_qty_payload(), facts, notes_seg=notes)
        self.assertEqual(card["method"], "量价绝对值法")

    def test_qty_identity_still_checked(self):
        facts = _facts()
        notes = _qty_notes()
        broken = _qty_payload()
        broken["historical_data"]["单价"]["2025A"] = 9000.0
        errors = validate_forecast(broken, facts, notes)
        self.assertTrue(any("对不上分部收入" in item for item in errors))

    def test_residual_forced_to_growth(self):
        facts = _facts()
        notes = {
            "name": "其他",
            "parent": "",
            "预测方法": "",
            "historical_revenue": {"2023A": 1.0, "2024A": 1.0, "2025A": 1.0},
        }
        qty_errors = validate_forecast(
            {
                "segment": "其他",
                "method": "量价绝对值法",
                "historical_data": {
                    "销量": {"2023A": 1.0, "2024A": 1.0, "2025A": 1.0},
                    "单价": {"2023A": 10000.0, "2024A": 10000.0, "2025A": 10000.0},
                },
                "final_forecast": {
                    "销量": {"2026E": 1.0, "2027E": 1.0, "2028E": 1.0},
                    "单价": {"2026E": 10000.0, "2027E": 10000.0, "2028E": 10000.0},
                },
                "unit_meta": {"fx": 10000.0},
                "final_rationale": "2026E残差按出货外推，这张卡不该用量价。",
            },
            facts,
            notes,
        )
        self.assertTrue(any("残差必须" in item for item in qty_errors))
        growth = {
            "segment": "其他",
            "method": "收入增速法",
            "historical_data": {"分部收入": {"2023A": 1.0, "2024A": 1.0, "2025A": 1.0}},
            "final_forecast": {"收入增速": {"2026E": 0.0, "2027E": 0.0, "2028E": 0.0}},
            "final_rationale": "2026E残差按0增速外推，收口项不跟量价。",
        }
        errors = validate_forecast(growth, facts, notes)
        self.assertFalse(any("残差必须" in item for item in errors))


class HistFillContractTests(unittest.TestCase):
    def test_parent_recon_and_no_parent_dump(self):
        names = ["消费电子", "汽车电子", "电脑互联", "通讯互联", "通讯及数据中心其他", "其他"]
        segs = [
            {"name": "消费电子", "parent": "消费电子"},
            {"name": "汽车电子", "parent": "汽车电子"},
            {"name": "电脑互联", "parent": "通讯及数据中心"},
            {"name": "通讯互联", "parent": "通讯及数据中心"},
            {"name": "通讯及数据中心其他", "parent": "通讯及数据中心"},
            {"name": "其他", "parent": ""},
        ]
        official_rows = [
            {"period": "2025A", "name": "消费电子", "revenue_yi": 700.0, "is_total": False},
            {"period": "2025A", "name": "汽车电子", "revenue_yi": 100.0, "is_total": False},
            {"period": "2025A", "name": "通讯及数据中心", "revenue_yi": 150.0, "is_total": False},
            {"period": "2025A", "name": "其他", "revenue_yi": 50.0, "is_total": False},
        ]
        deps = HistFillDeps(
            company="立讯精密",
            ticker="002475",
            hist_periods=["2025A"],
            forecast_years=["2026"],
            revenue=[1000.0],
            plan_names=names,
            plan_segments=segs,
            official_names=["消费电子", "汽车电子", "通讯及数据中心"],
            official_rows=official_rows,
            search=lambda *a, **k: [],
        )
        dumped = hist_errors(
            deps,
            _hist_draft(
                {
                    "消费电子": 700.0,
                    "汽车电子": 100.0,
                    "电脑互联": 0.0,
                    "通讯互联": 150.0,
                    "通讯及数据中心其他": 0.0,
                    "其他": 50.0,
                }
            ),
        )
        self.assertTrue(any("不要把父口径" in item for item in dumped))

        ok = hist_errors(
            deps,
            _hist_draft(
                {
                    "消费电子": 700.0,
                    "汽车电子": 100.0,
                    "电脑互联": 0.0,
                    "通讯互联": 0.0,
                    "通讯及数据中心其他": 150.0,
                    "其他": 50.0,
                }
            ),
        )
        self.assertEqual(ok, [])

    def test_missing_company_residual_fails_before_fill(self):
        deps = HistFillDeps(
            company="宁德时代",
            ticker="300750",
            hist_periods=["2025A"],
            forecast_years=["2026"],
            revenue=[4237.0183],
            plan_names=["动力电池系统", "储能电池系统"],
            plan_segments=[
                {"name": "动力电池系统", "parent": "动力电池系统"},
                {"name": "储能电池系统", "parent": "储能电池系统"},
            ],
            official_names=["动力电池系统", "储能电池系统"],
            official_rows=[
                {"period": "2025A", "name": "动力电池系统", "revenue_yi": 3165.0637, "is_total": False},
                {"period": "2025A", "name": "储能电池系统", "revenue_yi": 624.3982, "is_total": False},
            ],
            search=lambda *_a, **_k: [],
        )
        errors = plan_slot_errors(deps)
        self.assertTrue(any("公司级残差" in item for item in errors))
        deps.plan_names = ["动力电池系统", "储能电池系统", "其他"]
        deps.plan_segments.append({"name": "其他", "parent": ""})
        self.assertEqual(plan_slot_errors(deps), [])


class GrowthRateTests(unittest.TestCase):
    def test_percent_and_fraction(self):
        self.assertAlmostEqual(as_growth_rate(18.0), 18.0)
        self.assertAlmostEqual(as_growth_rate(0.18), 0.18)
        self.assertAlmostEqual(as_growth_rate(1.2), 1.2)

    def test_normalize_then_validate(self):
        facts = _facts()
        notes = _notes(10.0)
        errors = validate_forecast(_growth_payload(18.0, 10.0), facts, notes)
        self.assertTrue(any("不超过" in item for item in errors), errors)
        card = require_forecast(_growth_payload(0.18, 10.0), facts, notes_seg=notes)
        self.assertAlmostEqual(card["final_forecast"]["收入增速"]["2026E"], 0.18)
        self.assertEqual(validate_forecast(_growth_payload(1.2, 10.0), facts, notes), [])

    def test_zero_base_rejected(self):
        facts = _facts()
        notes = _notes(0.0)
        errors = validate_forecast(_growth_payload(0.12, 0.0), facts, notes)
        self.assertTrue(any("为 0" in item for item in errors))

    def test_residual_zero_base_allowed(self):
        facts = _facts()
        notes = {
            "name": "其他",
            "parent": "",
            "预测方法": "收入增速法",
            "historical_revenue": {"2023A": 1.0, "2024A": 1.0, "2025A": 0.0},
        }
        payload = {
            "segment": "其他",
            "method": "收入增速法",
            "historical_data": {"分部收入": {"2023A": 1.0, "2024A": 1.0, "2025A": 0.0}},
            "final_forecast": {"收入增速": {"2026E": 0.0, "2027E": 0.0, "2028E": 0.0}},
            "final_rationale": "2026E按0增速外推，去年残差已收口到0。",
        }
        errors = validate_forecast(payload, facts, notes)
        self.assertFalse(any("为 0" in item for item in errors))


class WorkbookTreeTests(unittest.TestCase):
    def test_parent_before_children_and_top_level_total(self):
        split = {
            "official_parents": ["消费电子", "汽车电子", "通讯及数据中心"],
            "final_segments": [
                {"name": "消费电子", "parent": "消费电子", "method": "收入增速法"},
                {"name": "汽车电子", "parent": "汽车电子", "method": "收入增速法"},
                {"name": "电脑互联", "parent": "通讯及数据中心", "method": "收入增速法"},
                {"name": "通讯互联", "parent": "通讯及数据中心", "method": "收入增速法"},
                {"name": "通讯及数据中心其他", "parent": "通讯及数据中心", "method": "收入增速法"},
                {"name": "其他", "parent": "", "method": "收入增速法"},
            ],
        }
        groups, company_res = build_split_groups(split)
        names = []
        for group in groups:
            names.append(group.top_name)
            names.extend(child["name"] for child in group.children)
        if company_res:
            names.append(company_res["name"])
        comms = names.index("通讯及数据中心")
        self.assertLess(comms, names.index("电脑互联"))
        self.assertLess(comms, names.index("通讯互联"))
        labels = top_level_revenue_labels(groups, company_res)
        self.assertEqual(
            labels,
            [
                "分部收入_消费电子",
                "分部收入_汽车电子",
                "分部收入_通讯及数据中心",
                "分部收入_其他",
            ],
        )
        self.assertNotIn("分部收入_电脑互联", labels)
        self.assertNotIn("分部收入_通讯互联", labels)


class SevenMethodContractTests(unittest.TestCase):
    def test_all_methods_identity_pass_and_fail(self):
        facts = _facts()
        notes = _qty_notes()
        for payload in (
            _qty_payload(),
            _qty_growth_payload(),
            _market_payload(),
            _order_payload(),
            _store_payload(),
            _user_payload(),
            _growth_payload(0.1, 10.0),
        ):
            card = require_forecast(payload, facts, notes_seg=notes)
            self.assertEqual(validate_forecast(card, facts, notes), [])
            broken = _break_identity(card)
            if broken is None:
                continue
            errors = validate_forecast(broken, facts, notes)
            self.assertTrue(any("对不上分部收入" in item for item in errors), errors)

    def test_fx_must_be_self_reported(self):
        facts = _facts()
        notes = _qty_notes()
        missing = _qty_payload()
        missing["unit_meta"] = {"volume_unit": "万只", "price_unit": "元"}
        errors = validate_forecast(missing, facts, notes)
        self.assertTrue(any("fx" in item for item in errors), errors)
        market = require_forecast(_market_payload(), facts, notes_seg=notes)
        self.assertEqual(validate_forecast(market, facts, notes), [])
        self.assertNotIn("fx", market.get("unit_meta") or {})

    def test_adapt_reads_method_from_card(self):
        facts = _facts()
        notes = {
            "company": "测",
            "ticker": "000000",
            "final_segments": [
                {
                    "name": "消费电子",
                    "parent": "消费电子",
                    "预测方法": "收入增速法",
                    "historical_revenue": {"2023A": 8.0, "2024A": 9.0, "2025A": 10.0},
                    "note": "",
                }
            ],
            "split_explanation": {
                "拆分逻辑": "测",
                "历史数据说明": "测",
                "其他业务说明": "测",
                "主要来源": "测",
            },
            "revenue_unit": "亿元",
        }
        adapted = adapt_split(notes, facts, {"消费电子": {"method": "量价绝对值法"}})
        self.assertEqual(adapted["final_segments"][0]["method"], "量价绝对值法")
        empty_notes = dict(notes)
        empty_notes["final_segments"] = [
            {**notes["final_segments"][0], "预测方法": ""}
        ]
        adapted_empty = adapt_split(empty_notes, facts, {"消费电子": {"预测方法": "市场渗透法"}})
        self.assertEqual(adapted_empty["final_segments"][0]["method"], "市场渗透法")

    def test_roll_does_not_read_revenue_growth(self):
        facts = _facts()
        hist = facts["hist_periods"]
        fcst = facts["forecast_periods"]
        qty = require_forecast(_qty_payload(), facts, notes_seg=_qty_notes())
        qty["final_forecast"]["收入增速"] = {"2026E": 0.99, "2027E": 0.99, "2028E": 0.99}
        revenue, drivers = roll_segment("量价绝对值法", qty, hist, fcst, _rev_map())
        self.assertAlmostEqual(revenue["2026E"], 12.0)
        self.assertNotIn("revenue_growth", drivers)
        market = require_forecast(_market_payload(), facts, notes_seg=_qty_notes())
        market["final_forecast"]["收入增速"] = {"2026E": 0.99, "2027E": 0.99, "2028E": 0.99}
        rolled = _segment_roll(
            {"name": "消费电子", "method": "市场渗透法", "hist_revenue": [8.0, 9.0, 10.0], "note": ""},
            market,
            hist,
            fcst,
        )
        self.assertAlmostEqual(rolled["revenue"]["2026E"], 11.0)
        self.assertNotIn("revenue_growth", rolled["drivers"])

    def test_residual_must_use_growth_method(self):
        facts = _facts()
        notes = {
            "name": "通讯及数据中心其他",
            "parent": "通讯及数据中心",
            "预测方法": "",
            "historical_revenue": {"2023A": 8.0, "2024A": 9.0, "2025A": 10.0},
            "data_quality": {"2023A": "兜底", "2024A": "兜底", "2025A": "兜底"},
            "note": "[兜底] 父口径收口",
        }
        self.assertTrue(is_residual_name("通讯及数据中心其他", parent="通讯及数据中心"))
        qty_errors = validate_forecast(_qty_payload() | {"segment": "通讯及数据中心其他"}, facts, notes)
        self.assertTrue(any("残差必须" in item for item in qty_errors))
        growth = _growth_payload(0.0, 10.0)
        growth["segment"] = "通讯及数据中心其他"
        growth["historical_data"]["分部收入"] = notes["historical_revenue"]
        self.assertEqual(validate_forecast(growth, facts, notes), [])
        self.assertEqual(_note_tag("有据可查"), "[有据可查]")
        self.assertEqual(_note_tag("兜底"), "[兜底]")

    def test_certify_uses_t_recon(self):
        self.assertGreater(t_recon(100.0), 0.05)
        facts = {
            "hist_periods": ["2025A"],
            "income": {
                "营业收入": [100.0],
                "税前利润": [20.0],
                "营业利润": [19.94],
                "营业外收入": [0.0],
                "营业外支出": [0.0],
            },
        }
        segs = [{"hist_revenue": [100.06]}]
        self.assertEqual(hist_recon_errors(facts, segs), [])
        segs_bad = [{"hist_revenue": [101.0]}]
        fails = hist_recon_errors(facts, segs_bad)
        self.assertTrue(any("分部加总" in item for item in fails))

    def test_formula_tokens(self):
        self.assertEqual(formula_errors("量价绝对值法", "=A10*B10/单位换算因子_消费电子"), [])
        self.assertEqual(formula_errors("量价绝对值法", "=E24*E25/E5"), [])
        self.assertTrue(formula_errors("量价绝对值法", "=A10*B10"))
        self.assertEqual(formula_errors("市场渗透法", "=A10*B10*C10"), [])
        self.assertEqual(formula_errors("订单转化法", "=(A10+B10)*C10"), [])
        self.assertEqual(formula_errors("收入增速法", "=C9*(1+D9)"), [])
        self.assertTrue(formula_errors("收入增速法", "=A10*B10/单位换算因子_x"))
        self.assertTrue(formula_errors("量价增速法", "=C9*(1+D9)"))

    def test_formula_row_labels_resolve_fx(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws["A5"] = "单位换算因子"
        ws["A24"] = "销量"
        ws["A25"] = "单价"
        self.assertEqual(
            formula_row_labels(ws, "=E24*E25/E5"),
            ["销量", "单价", "单位换算因子"],
        )

    def test_find_row_in_group_does_not_cross_segment(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws["A2"] = "收入驱动假设区"
        ws["A3"] = "动力电池系统（量价增速法）"
        ws["A4"] = "单位换算因子"
        ws["A5"] = "销量增速"
        ws["A6"] = "储能电池系统（量价增速法）"
        ws["A7"] = "单位换算因子"
        ws["A8"] = "收入构建区"
        ws["A9"] = "动力电池系统（量价增速法）"
        ws["A10"] = "销量"
        ws["A11"] = "分部收入"
        ws["A12"] = "储能电池系统（量价增速法）"
        ws["A13"] = "销量"
        ws["A14"] = "分部收入"
        ws["A15"] = "收入归因"
        titles = {
            "动力电池系统（量价增速法）",
            "储能电池系统（量价增速法）",
        }
        self.assertEqual(
            find_row_in_group(
                ws,
                "动力电池系统（量价增速法）",
                "分部收入",
                after_row=8,
                until_row=15,
                group_titles=titles,
            ),
            11,
        )
        self.assertEqual(
            find_row_in_group(
                ws,
                "储能电池系统（量价增速法）",
                "单位换算因子",
                after_row=2,
                until_row=8,
                group_titles=titles,
            ),
            7,
        )


def _hist_draft(amounts: dict[str, float]) -> HistFill:
    return HistFill(
        fill_method="残差倒推",
        segments=[
            HistSeg(
                name=name,
                years=[HistCell(year="2025A", amount=amount, quality="倒推", why="测")],
            )
            for name, amount in amounts.items()
        ],
        split_logic="测",
        hist_data_note="测",
        other_note="测",
        sources_text="测",
        final_rationale="测",
    )


def _facts() -> dict:
    return {
        "hist_periods": ["2023A", "2024A", "2025A"],
        "forecast_periods": ["2026E", "2027E", "2028E"],
        "unit": "亿元",
        "income": {"营业收入": [100.0, 110.0, 120.0]},
        "as_of": "2026-09-01",
    }


def _notes(last: float) -> dict:
    return {
        "name": "消费电子",
        "预测方法": "收入增速法",
        "historical_revenue": {"2023A": 8.0, "2024A": 9.0, "2025A": last},
    }


def _growth_payload(growth: float, last: float) -> dict:
    return {
        "segment": "消费电子",
        "method": "收入增速法",
        "historical_data": {"分部收入": {"2023A": 8.0, "2024A": 9.0, "2025A": last}},
        "final_forecast": {"收入增速": {"2026E": growth, "2027E": 0.1, "2028E": 0.1}},
        "final_rationale": "2026E我们按事件发酵给出增速，客户放量带动结构改善。",
    }


def _qty_notes() -> dict:
    return {
        "name": "消费电子",
        "parent": "消费电子",
        "预测方法": "",
        "historical_revenue": {"2023A": 8.0, "2024A": 9.0, "2025A": 10.0},
    }


def _qty_payload() -> dict:
    return {
        "segment": "消费电子",
        "method": "量价绝对值法",
        "historical_data": {
            "销量": {"2023A": 8.0, "2024A": 9.0, "2025A": 10.0},
            "单价": {"2023A": 10000.0, "2024A": 10000.0, "2025A": 10000.0},
        },
        "final_forecast": {
            "销量": {"2026E": 12.0, "2027E": 13.0, "2028E": 14.0},
            "单价": {"2026E": 10000.0, "2027E": 10000.0, "2028E": 10000.0},
        },
        "unit_meta": {"fx": 10000.0},
        "final_rationale": "2026E出货12，单价维持，收入对得上年报锚。",
    }


def _qty_growth_payload() -> dict:
    return {
        "segment": "消费电子",
        "method": "量价增速法",
        "historical_data": {
            "销量": {"2023A": 8.0, "2024A": 9.0, "2025A": 10.0},
            "单价": {"2023A": 10000.0, "2024A": 10000.0, "2025A": 10000.0},
        },
        "final_forecast": {
            "销量增速": {"2026E": 0.1, "2027E": 0.1, "2028E": 0.1},
            "单价增速": {"2026E": 0.0, "2027E": 0.0, "2028E": 0.0},
        },
        "unit_meta": {"fx": 10000.0},
        "final_rationale": "2026E按出货增速外推，单价持平，收入对得上年报锚。",
    }


def _market_payload() -> dict:
    return {
        "segment": "消费电子",
        "method": "市场渗透法",
        "historical_data": {
            "市场规模": {"2023A": 80.0, "2024A": 90.0, "2025A": 100.0},
            "渗透率": {"2023A": 0.2, "2024A": 0.2, "2025A": 0.2},
            "市占率": {"2023A": 0.5, "2024A": 0.5, "2025A": 0.5},
        },
        "final_forecast": {
            "市场规模": {"2026E": 110.0, "2027E": 120.0, "2028E": 130.0},
            "渗透率": {"2026E": 0.2, "2027E": 0.2, "2028E": 0.2},
            "市占率": {"2026E": 0.5, "2027E": 0.5, "2028E": 0.5},
        },
        "final_rationale": "2026E按市场规模和份额拆，收入对得上年报锚。",
    }


def _order_payload() -> dict:
    return {
        "segment": "消费电子",
        "method": "订单转化法",
        "historical_data": {
            "期初在手订单": {"2023A": 6.0, "2024A": 7.0, "2025A": 8.0},
            "新签订单": {"2023A": 10.0, "2024A": 11.0, "2025A": 12.0},
            "转化率": {"2023A": 0.5, "2024A": 0.5, "2025A": 0.5},
        },
        "final_forecast": {
            "期初在手订单": {"2026E": 9.0, "2027E": 10.0, "2028E": 11.0},
            "新签订单": {"2026E": 13.0, "2027E": 14.0, "2028E": 15.0},
            "转化率": {"2026E": 0.5, "2027E": 0.5, "2028E": 0.5},
        },
        "final_rationale": "2026E按在手加新签转化，收入对得上年报锚。",
    }


def _store_payload() -> dict:
    return {
        "segment": "消费电子",
        "method": "门店坪效法",
        "historical_data": {
            "门店数": {"2023A": 8.0, "2024A": 9.0, "2025A": 10.0},
            "单店产出": {"2023A": 10000.0, "2024A": 10000.0, "2025A": 10000.0},
        },
        "final_forecast": {
            "门店数": {"2026E": 11.0, "2027E": 12.0, "2028E": 13.0},
            "单店产出": {"2026E": 10000.0, "2027E": 10000.0, "2028E": 10000.0},
        },
        "unit_meta": {"fx": 10000.0, "store_unit": "家", "productivity_unit": "万元"},
        "final_rationale": "2026E按门店和单店产出拆，收入对得上年报锚。",
    }


def _user_payload() -> dict:
    return {
        "segment": "消费电子",
        "method": "用户单价法",
        "historical_data": {
            "订阅用户": {"2023A": 8.0, "2024A": 9.0, "2025A": 10.0},
            "单用户收入": {"2023A": 10000.0, "2024A": 10000.0, "2025A": 10000.0},
        },
        "final_forecast": {
            "订阅用户": {"2026E": 11.0, "2027E": 12.0, "2028E": 13.0},
            "单用户收入": {"2026E": 10000.0, "2027E": 10000.0, "2028E": 10000.0},
        },
        "unit_meta": {"fx": 10000.0, "volume_unit": "万户", "price_unit": "元"},
        "final_rationale": "2026E按用户和ARPU拆，收入对得上年报锚。",
    }


def _rev_map() -> dict[str, float]:
    return {"2023A": 8.0, "2024A": 9.0, "2025A": 10.0}


def _break_identity(card: dict) -> dict | None:
    method = card["method"]
    if method == "收入增速法":
        return None
    hist = dict(card["historical_data"])
    field = {
        "量价绝对值法": "单价",
        "量价增速法": "单价",
        "市场渗透法": "市占率",
        "订单转化法": "转化率",
        "门店坪效法": "单店产出",
        "用户单价法": "单用户收入",
    }[method]
    series = dict(hist[field])
    series["2025A"] = float(series["2025A"]) * 1.5
    hist[field] = series
    return {**card, "historical_data": hist}


if __name__ == "__main__":
    unittest.main()
